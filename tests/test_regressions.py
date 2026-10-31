"""Run with: python -m unittest discover -s tests -v. No real data or email used."""
import os
import re
import io
import sys
import time
import sqlite3
import tempfile
import subprocess
import json
import html
import unittest
from pathlib import Path
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_boot = tempfile.TemporaryDirectory()
os.environ.update(CLASS_MANAGER_DB=str(Path(_boot.name)/'boot.db'), CLASS_MANAGER_INSTANCE=str(Path(_boot.name)/'instance'),
                  CLASS_MANAGER_BACKUPS=str(Path(_boot.name)/'backups'), CLASS_MANAGER_ENV_FILE=str(Path(_boot.name)/'no-env'),
                  CLASS_MANAGER_ENV='test', SECRET_KEY='test-secret-key-with-more-than-thirty-two-characters', ADMIN_INITIAL_PASSWORD='InitialPass123!')
import app as site
from database import init_database, backup_database
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook, Workbook
from flask import template_rendered


class SiteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db = site.DB
        site.DB = str(Path(self.tmp.name)/'test.db')
        init_database(site.DB, Path(self.tmp.name)/'backups', Path(self.tmp.name)/'instance', 'InitialPass123!')
        site.app.config.update(TESTING=True)
        self.sql('UPDATE users SET must_change_password=0')
        for name in ['teacher_a', 'teacher_b']:
            self.sql('INSERT INTO users(username,password,email) VALUES(?,?,?)', (name, generate_password_hash('TeacherPass123!'), name+'@example.invalid'))
        self.a = self.client('teacher_a')
        self.b = self.client('teacher_b')
        self.admin = self.client('admin', 'InitialPass123!')
        self.post(self.a, '/add', {'name':'学生甲'})
        self.post(self.a, '/add', {'name':'学生乙'})
        self.post(self.b, '/add', {'name':'其他老师学生'})
        self.sa,self.s2,self.sb = [r['id'] for r in self.sql('SELECT id FROM students ORDER BY id')]
        self.post(self.a, '/create_exam', {'name':'期中考试'})
        self.post(self.b, '/create_exam', {'name':'期中考试'})
        self.ea,self.eb = [r['id'] for r in self.sql('SELECT id FROM exams ORDER BY id')]
        self.ca = self.sql("SELECT id FROM exam_subjects WHERE exam_id=? AND subject='语文'", (self.ea,))[0][0]
        self.cb = self.sql("SELECT id FROM exam_subjects WHERE exam_id=? AND subject='语文'", (self.eb,))[0][0]

    def tearDown(self):
        site.DB = self.old_db
        self.tmp.cleanup()

    def sql(self, query, args=()):
        with site.database(write=True) as conn:
            return conn.execute(query, args).fetchall()

    def post(self, client, path, data=None, follow=False):
        client.get('/login', follow_redirects=True)
        with client.session_transaction() as sess:
            token = sess['csrf_token']
        return client.post(path, data={**(data or {}), 'csrf_token':token}, follow_redirects=follow)

    def client(self, name=None, password='TeacherPass123!'):
        client = site.app.test_client()
        if name:
            user = self.sql('SELECT email FROM users WHERE username=?', (name,))[0]
            response = self.post(client, '/login', {'email':user['email'] or name,'password':password})
            self.assertEqual(response.status_code,302)
        return client

    def save(self, sid=None, exam=None, subject='语文', score='90', status='recorded', client=None):
        return self.post(client or self.a, f'/score/{sid or self.sa}', {'exam_id':exam or self.ea,'subject':subject,'score':score,'status':status})

    def summary(self):
        contexts=[]
        def capture(sender,template,context,**kw):contexts.append(context)
        template_rendered.connect(capture,site.app)
        try:self.a.get(f'/summary?exam_id={self.ea}')
        finally:template_rendered.disconnect(capture,site.app)
        return contexts[-1]

    def test_all_pages_and_auth_guards(self):
        anonymous = self.client()
        for path in ['/', '/exams','/summary','/export_excel','/help','/change_password',f'/exam/{self.ea}',f'/score/{self.sa}',f'/student/{self.sa}']:
            self.assertEqual(anonymous.get(path).status_code,302,path)
            self.assertEqual(self.a.get(path).status_code,200,path)
        self.assertEqual(self.a.get('/users').status_code,403)
        self.assertEqual(self.admin.get('/users').status_code,200)

    def test_csrf_blocks_missing_and_wrong_tokens(self):
        for token in [None,'wrong']:
            data={'name':'CSRF'}
            if token:data['csrf_token']=token
            self.assertEqual(self.a.post('/add',data=data).status_code,400)
        self.assertFalse(self.sql('SELECT 1 FROM students WHERE name=?',('CSRF',)))
        self.assertEqual(self.client().post('/login',data={'username':'admin','password':'InitialPass123!'}).status_code,400)

    def test_forged_old_secret_rejected(self):
        from flask import Flask
        fake=Flask('fake');fake.secret_key='student-grade-system-secret-2026'
        token=fake.session_interface.get_signing_serializer(fake).dumps({'user':'admin','uid':1,'session_version':1})
        client=self.client();client.set_cookie('session',token)
        self.assertEqual(client.get('/users').status_code,302)

    def test_student_and_exam_isolation(self):
        self.assertNotIn('其他老师学生',self.a.get('/').text)
        for path in [f'/student/{self.sb}',f'/score/{self.sb}',f'/exam/{self.eb}',f'/summary?exam_id={self.eb}']:
            self.assertEqual(self.a.get(path).status_code,404,path)
        self.assertEqual(self.post(self.a,f'/delete/{self.sb}').status_code,404)
        self.assertEqual(self.post(self.a,f'/delete_exam/{self.eb}').status_code,404)

    def test_cross_teacher_exam_write_and_read(self):
        self.assertEqual(self.save(exam=self.eb).status_code,404)
        self.assertEqual(self.a.get(f'/score/{self.sa}?exam_id={self.eb}').status_code,404)
        self.assertFalse(self.sql('SELECT 1 FROM scores'))
        self.assertEqual(self.save(exam=self.eb,client=self.admin).status_code,400)

    def test_subject_ids_must_match_exam(self):
        self.assertEqual(self.post(self.a,f'/exam/{self.ea}/update_subject/{self.cb}',{'max_score':123,'pass_score':70}).status_code,404)
        self.assertEqual(self.post(self.a,f'/exam/{self.ea}/delete_subject/{self.cb}').status_code,404)
        self.assertEqual(self.sql('SELECT max_score FROM exam_subjects WHERE id=?',(self.cb,))[0][0],150)

    def test_subject_tamper_does_not_leave_scores(self):
        self.save()
        self.post(self.a,f'/exam/{self.ea}/delete_subject/{self.ca}',{'subject':'其他名称'})
        self.assertFalse(self.sql('SELECT 1 FROM scores'))
        self.assertEqual(self.a.get(f'/student/{self.sa}').status_code,200)

    def test_normal_save_overwrite_edit_and_delete(self):
        self.assertEqual(self.save(score='90').status_code,302)
        self.save(score='88.5')
        rows=self.sql('SELECT * FROM scores')
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['score'],88.5)
        self.assertEqual(self.post(self.a,f'/edit_score/{rows[0]["id"]}',{'score':99}).status_code,302)
        self.assertEqual(self.sql('SELECT score FROM scores')[0][0],99)
        self.assertEqual(self.post(self.b,f'/edit_score/{rows[0]["id"]}',{'score':100}).status_code,404)
        self.assertEqual(self.post(self.b,f'/delete_score/{rows[0]["id"]}').status_code,404)
        self.post(self.a,f'/delete_score/{rows[0]["id"]}')
        self.assertFalse(self.sql('SELECT 1 FROM scores'))

    def test_invalid_numbers_in_save_and_edit(self):
        for value in ['nan','NaN','inf','-inf','1e309','1e308','abc','-1','151','']:
            self.save(score=value)
            self.assertFalse(self.sql('SELECT 1 FROM scores'),value)
        self.save(score=90);sid=self.sql('SELECT id FROM scores')[0][0]
        for value in ['nan','inf','-1','151']:
            self.post(self.a,f'/edit_score/{sid}',{'score':value})
            self.assertEqual(self.sql('SELECT score FROM scores')[0][0],90)

    def test_invalid_subject_numbers_and_max_reduction(self):
        for value in ['nan','inf','-1','0','abc']:
            self.post(self.a,f'/exam/{self.ea}/add_subject',{'subject':'无效','max_score':value,'pass_score':60})
        self.assertFalse(self.sql('SELECT 1 FROM exam_subjects WHERE subject=?',('无效',)))
        self.save(score=149)
        self.post(self.a,f'/exam/{self.ea}/update_subject/{self.ca}',{'max_score':100,'pass_score':60})
        self.assertEqual(self.sql('SELECT max_score FROM exam_subjects WHERE id=?',(self.ca,))[0][0],150)
        for value in ['nan','inf']:
            self.post(self.a,f'/exam/{self.ea}/update_subject/{self.ca}',{'max_score':150,'pass_score':value})
            self.assertEqual(self.sql('SELECT pass_score FROM exam_subjects WHERE id=?',(self.ca,))[0][0],90)

    def test_absent_is_distinct_from_zero_and_missing(self):
        self.save(status='absent',score='')
        self.assertIsNone(self.sql('SELECT score FROM scores')[0][0])
        context=self.summary()
        a=next(r for r in context['incomplete'] if r['sid']==self.sa)
        self.assertEqual((a['absent'],a['missing'],a['count']),(1,5,0))
        self.save(score=0)
        a=next(r for r in self.summary()['incomplete'] if r['sid']==self.sa)
        self.assertEqual((a['absent'],a['missing'],a['fail_count']),(0,5,1))

    def test_ranking_excludes_partial_and_ties(self):
        # Restrict this exam to two subjects, then test ranks 1,1,3.
        self.sql("DELETE FROM exam_subjects WHERE exam_id=? AND subject NOT IN ('语文','数学')",(self.ea,))
        self.save(score=150)
        self.save(sid=self.s2,score=149);self.save(sid=self.s2,subject='数学',score=149)
        context=self.summary();self.assertEqual([r['sid'] for r in context['ranking']],[self.s2])
        self.save(subject='数学',score=148)
        self.post(self.a,'/add',{'name':'第三人'});third=self.sql('SELECT id FROM students WHERE name=?',('第三人',))[0][0]
        self.save(sid=third,score=100);self.save(sid=third,subject='数学',score=100)
        self.assertEqual([r['rank'] for r in self.summary()['ranking']],[1,1,3])

    def test_empty_summary_no_all_pass_claim(self):
        self.assertNotIn('全员及格',self.a.get('/summary').text)
        self.save(score=100)
        self.assertIn('暂不能判断全员及格',self.a.get('/summary').text)

    def test_statistics_use_valid_denominator(self):
        self.save(score=90);self.save(sid=self.s2,score=60)
        s=next(s for s in self.summary()['statistics'] if s['subject']=='语文')
        self.assertEqual((s['average'],s['pass_rate'],s['excellent_rate'],s['count']),(75,50,0,2))

    def test_dynamic_excel_and_text_names(self):
        self.post(self.a,f'/exam/{self.ea}/add_subject',{'subject':'历史','max_score':100,'pass_score':60})
        self.save(subject='历史',score=59)
        self.post(self.a,'/add',{'name':'=1+1'})
        response=self.a.get('/export_excel');self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data));sheet=book.active
        headers=[c.value for c in sheet[1]];col=headers.index('历史')+1
        self.assertEqual(sheet.cell(2,col).value,59)
        self.assertEqual(sheet.cell(2,col).fill.fgColor.rgb,'00FECACA')
        cell=next(row[0] for row in sheet.iter_rows(min_row=2) if row[0].value=='=1+1')
        self.assertEqual(cell.data_type,'s')
        self.assertNotIn('其他老师学生',[r[0].value for r in sheet.iter_rows()])

    def test_excel_sheet_names_invalid_and_collision(self):
        for name in ['期中/测试','期中\\测试','[]:*?','x'*40+'1','x'*40+'2',"'引号'"]:
            self.post(self.a,'/create_exam',{'name':name})
        response=self.a.get('/export_excel');self.assertEqual(response.status_code,200)
        names=load_workbook(io.BytesIO(response.data)).sheetnames
        self.assertEqual(len(names),7);self.assertEqual(len({n.casefold() for n in names}),7)
        self.assertTrue(all(len(n)<=31 and not re.search(r'[\\/*?:\[\]]',n) for n in names))

    def test_xss_safe_messages_and_confirmations(self):
        name="</script><script>window.audit=1</script>"
        self.post(self.a,'/add',{'name':name})
        response=self.post(self.a,'/add',{'name':name},follow=True)
        self.assertNotIn(name,response.text)
        self.assertIn('&lt;script&gt;',response.text)
        self.assertNotIn("confirm('确定删除 "+name,response.text)

    def test_deleted_account_sessions_and_data(self):
        self.save()
        uid=self.sql('SELECT id FROM users WHERE username=?',('teacher_a',))[0][0]
        self.post(self.admin,f'/delete_user/{uid}')
        self.assertEqual(self.a.get('/').status_code,302)
        self.assertEqual(self.post(self.client(),'/login',{'username':'teacher_a','password':'TeacherPass123!'}).status_code,200)
        self.assertTrue(self.sql('SELECT 1 FROM scores'))
        self.post(self.admin,f'/enable_user/{uid}')
        self.client('teacher_a')
        self.assertEqual(self.a.get('/').status_code,302)
        self.post(self.admin,'/delete_user/1')
        self.assertEqual(self.sql('SELECT active FROM users WHERE id=1')[0][0],1)

    def test_password_change_revokes_other_sessions(self):
        other=self.client('teacher_a')
        self.post(self.a,'/change_password',{'old_password':'TeacherPass123!','new_password':'ChangedPass123!','confirm_password':'ChangedPass123!'})
        self.assertEqual(self.a.get('/').status_code,200)
        self.assertEqual(other.get('/').status_code,302)
        self.assertEqual(self.a.get('/logout').status_code,405)
        self.post(self.a,'/logout')
        self.assertEqual(self.a.get('/').status_code,302)

    def reset_code(self, client=None):
        calls=[]
        with patch.object(site,'send_email',side_effect=lambda *args:calls.append(args)):
            response=self.post(client or self.a,'/forgot_password',{'email':'teacher_a@example.invalid'})
        self.assertEqual(response.status_code,302)
        return re.search(r'验证码：(\d{6})',calls[0][2]).group(1)

    def reset(self, code, client=None, **extra):
        return self.post(client or self.a,'/reset_password',dict(email='teacher_a@example.invalid',token=code,new_password='NewPassword123!',confirm='NewPassword123!',**extra))

    def test_reset_hashed_token_single_use_and_sessions(self):
        other=self.client('teacher_a');code=self.reset_code()
        self.assertNotEqual(self.sql('SELECT token FROM reset_tokens')[0][0],code)
        self.assertEqual(self.reset(code).status_code,302)
        self.assertFalse(self.sql('SELECT 1 FROM reset_tokens'))
        self.assertEqual(other.get('/').status_code,302)
        self.assertEqual(self.reset(code).status_code,200)
        self.client('teacher_a','NewPassword123!')

    def test_reset_attempt_limit_and_expiry(self):
        code=self.reset_code();wrong='000000' if code!='000000' else '111111'
        for _ in range(5):self.reset(wrong)
        self.assertEqual(self.reset(code).status_code,200)
        self.assertTrue(self.sql('SELECT 1 FROM reset_tokens'))
        self.sql('UPDATE reset_tokens SET attempts=0,expires_at=?',(time.time()-1,))
        self.assertEqual(self.reset(code).status_code,200)

    def test_send_limits_and_failed_delivery_cleanup(self):
        with patch.object(site,'send_email',side_effect=RuntimeError('provider-secret-should-not-leak')) as send:
            response=self.post(self.a,'/forgot_password',{'email':'teacher_a@example.invalid'})
            self.assertNotIn('provider-secret',response.text)
            self.assertFalse(self.sql('SELECT 1 FROM reset_tokens'))
            response=self.post(self.a,'/forgot_password',{'email':'teacher_a@example.invalid'})
            self.assertEqual(response.status_code,429);self.assertEqual(send.call_count,1)

    def test_expired_token_cleared_and_unknown_email_generic(self):
        self.sql('INSERT INTO reset_tokens(email,token,expires_at) VALUES(?,?,?)',('old@example.invalid','bad',0))
        with patch.object(site,'send_email') as send:
            response=self.post(self.a,'/forgot_password',{'email':'unknown@example.invalid'})
            self.assertEqual(response.status_code,302);send.assert_not_called()
        self.assertFalse(self.sql('SELECT 1 FROM reset_tokens'))
        self.assertIn('如果该邮箱已注册',self.a.get(response.location).text)

    def test_login_limit(self):
        c=self.client()
        for _ in range(10):self.post(c,'/login',{'username':'bad','password':'bad'})
        self.assertEqual(self.post(c,'/login',{'username':'bad','password':'bad'}).status_code,429)

    def test_registration_and_duplicate_case_insensitive_email(self):
        c=self.client()
        data=dict(username='new_teacher',password='TeacherPass123!',confirm='TeacherPass123!',email='NEW@example.invalid')
        self.assertEqual(self.post(c,'/register',data).status_code,302)
        self.assertIn('用户名已被占用',self.post(c,'/register',data).text)
        self.assertIn('该邮箱已被注册',self.post(c,'/register',{**data,'username':'other','email':'new@example.invalid'}).text)
        self.assertIn('密码长度',self.post(c,'/register',{**data,'password':'1'}).text)

    def test_email_login_and_legacy_username_fallback(self):
        c=self.client()
        response=self.post(c,'/login',{'email':'TEACHER_A@example.invalid','password':'TeacherPass123!'})
        self.assertEqual(response.status_code,302)
        self.sql("UPDATE users SET email=NULL WHERE username='teacher_b'")
        legacy=self.client()
        response=self.post(legacy,'/login',{'email':'teacher_b','password':'TeacherPass123!'})
        self.assertEqual(response.status_code,302)
        self.assertIn('尚未登记邮箱',legacy.get('/').text)

    def test_multiple_classes_isolate_students_and_exams(self):
        response=self.post(self.a,'/create_class',{'name':'二班','grade':'七年级'})
        self.assertEqual(response.status_code,302)
        second=self.sql("SELECT id FROM classes WHERE teacher_id=(SELECT id FROM users WHERE username='teacher_a') AND name='二班'")[0][0]
        self.assertEqual(self.a.get('/classes').status_code,200)
        self.post(self.a,'/add',{'name':'二班学生','class_id':second})
        self.post(self.a,'/create_exam',{'name':'二班期中','class_id':second})
        self.assertIn('二班学生',self.a.get(f'/?class_id={second}').text)
        self.assertNotIn('二班学生',self.a.get('/').text)
        self.assertIn('二班期中',self.a.get(f'/exams?class_id={second}').text)
        self.assertNotIn('二班期中',self.a.get('/exams').text)

    def test_student_code_unique_within_class(self):
        self.assertEqual(self.post(self.a,'/add',{'name':'学号学生','student_code':'A001'}).status_code,302)
        self.post(self.a,'/add',{'name':'重复学号','student_code':'A001'})
        self.assertEqual(len(self.sql("SELECT id FROM students WHERE class_id=? AND student_code='A001'", (self.sql("SELECT id FROM classes WHERE teacher_id=(SELECT id FROM users WHERE username='teacher_a') ORDER BY id LIMIT 1")[0][0],))),1)
        self.post(self.a,'/create_class',{'name':'学号测试二班'})
        cid=self.sql("SELECT id FROM classes WHERE name='学号测试二班'")[0][0]
        self.assertEqual(self.post(self.a,'/add',{'name':'另一班学生','student_code':'A001','class_id':cid}).status_code,302)

    def test_edit_student_profile_and_move_class(self):
        sid = self.sa
        self.post(self.a, f'/edit_student/{sid}', {'name':'学生甲改名','student_code':'A100','class_id':self.sql("SELECT id FROM classes WHERE teacher_id=(SELECT id FROM users WHERE username='teacher_a') ORDER BY id LIMIT 1")[0][0]})
        row=self.sql('SELECT name,student_code FROM students WHERE id=?',(sid,))[0]
        self.assertEqual((row['name'],row['student_code']),('学生甲改名','A100'))
        self.assertEqual(self.b.get(f'/edit_student/{sid}').status_code,404)

    def roster_file(self, rows, headers=('姓名', '学号')):
        wb = Workbook()
        wb.active.append(list(headers))
        for row in rows:
            wb.active.append(list(row))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def roster_preview(self, rows, client=None, cid=None, headers=('姓名', '学号')):
        cid = cid or self.sql('SELECT class_id FROM students WHERE id=?', (self.sa,))[0][0]
        return self.post(client or self.a, '/students/import',
                         {'class_id':cid, 'file':(self.roster_file(rows, headers), 'roster.xlsx')})

    def preview_token(self, response):
        match = re.search(r'name="preview_token" value="([^"]+)"', response.text)
        return match[1] if match else None

    def test_import_preview_then_atomic_confirm_and_replay(self):
        page = self.roster_preview([('新学生甲','001'), ('新学生乙','002')])
        self.assertEqual(page.status_code, 200)
        self.assertFalse(self.sql("SELECT 1 FROM students WHERE name='新学生甲'"))
        token = self.preview_token(page)
        self.assertTrue(token)
        result = self.post(self.a, '/students/import', {'action':'confirm','preview_token':token})
        self.assertEqual(result.status_code, 302)
        created = self.sql("SELECT * FROM students WHERE name LIKE '新学生%' ORDER BY id")
        self.assertEqual([r['student_code'] for r in created], ['001','002'])
        original = self.sql('SELECT teacher_id,class_id FROM students WHERE id=?', (self.sa,))[0]
        self.assertTrue(all((r['teacher_id'],r['class_id']) == tuple(original) for r in created))
        again = self.post(self.a, '/students/import', {'action':'confirm','preview_token':token})
        self.assertIn('已导入', again.text)
        self.assertEqual(len(self.sql("SELECT id FROM students WHERE name LIKE '新学生%'")), 2)

    def test_import_all_error_rows_reported_without_writing(self):
        self.sql('UPDATE students SET student_code=? WHERE id=?', ('001',self.sa))
        page = self.roster_preview([('学生甲','001'), ('新学生','002'), ('新学生','002'), ('','003'), ('超长学号','x'*51), ('=1+1','004')])
        for message in ['已有同名','已有相同学号','与第3行重复','姓名需为','学号不能超过','不能使用公式']:
            self.assertIn(message, page.text)
        self.assertIsNone(self.preview_token(page))
        self.assertEqual(len(self.sql('SELECT id FROM students')), 3)

    def test_import_rechecks_roster_changes_before_confirm(self):
        token = self.preview_token(self.roster_preview([('待导入甲','010'),('待导入乙','011')]))
        self.post(self.a, '/add', {'name':'另一个人','student_code':'011'})
        page = self.post(self.a, '/students/import', {'action':'confirm','preview_token':token})
        self.assertIn('名单已变化', page.text)
        self.assertFalse(self.sql("SELECT 1 FROM students WHERE name='待导入甲'"))

    def test_import_permission_token_tamper_expiry_and_changed_session(self):
        cid = self.sql('SELECT class_id FROM students WHERE id=?', (self.sa,))[0][0]
        self.assertEqual(self.b.get(f'/students/import?class_id={cid}').status_code, 404)
        self.assertEqual(self.roster_preview([('无权限','001')], client=self.b, cid=cid).status_code, 404)
        token = self.preview_token(self.roster_preview([('隔离测试','001')]))
        data = {'action':'confirm','preview_token':token}
        self.assertEqual(self.post(self.b, '/students/import', data).status_code, 403)
        self.assertEqual(self.post(self.a, '/students/import', {**data,'preview_token':token+'x'}).status_code, 400)
        from itsdangerous import URLSafeTimedSerializer
        signer = URLSafeTimedSerializer(site.app.secret_key, salt='student-import-v1')
        payload = signer.loads(token)
        with patch('itsdangerous.timed.time.time', return_value=time.time()-1801):
            expired = signer.dumps(payload)
        self.assertEqual(self.post(self.a, '/students/import', {**data,'preview_token':expired}).status_code, 400)
        self.post(self.a, '/change_password', {'old_password':'TeacherPass123!',
                  'new_password':'ChangedTeacher123!','confirm_password':'ChangedTeacher123!'})
        self.assertEqual(self.post(self.a, '/students/import', data).status_code, 403)
        self.assertFalse(self.sql("SELECT 1 FROM students WHERE name='隔离测试'"))

    def test_import_template_formats_codes_as_text(self):
        response = self.a.get('/students/import-template')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual([c.value for c in wb.active[1]], ['姓名','学号'])
        self.assertEqual(wb.active['B2'].number_format, '@')
        self.assertIsNone(wb.active['A2'].value)

    def test_import_class_binding_and_duplicate_codes_in_other_class(self):
        original_class = self.sql('SELECT class_id FROM students WHERE id=?', (self.sa,))[0][0]
        self.post(self.a, '/create_class', {'name':'导入二班'})
        other_class = self.sql("SELECT id FROM classes WHERE name='导入二班'")[0][0]
        self.post(self.a, '/add', {'name':'相同姓名','student_code':'001','class_id':other_class})
        token = self.preview_token(self.roster_preview([('相同姓名','001')]))
        self.assertTrue(token)
        page = self.post(self.a, '/students/import',
                         {'action':'confirm','preview_token':token,'class_id':other_class})
        self.assertEqual(page.status_code, 302)
        rows = self.sql("SELECT class_id FROM students WHERE name='相同姓名'")
        self.assertEqual({r[0] for r in rows}, {original_class,other_class})

    def test_import_file_headers_limits_and_corrupt_files(self):
        for rows, headers, expected in [([], ('姓名','学号'), '没有可导入'),
                                       ([('甲','1')], ('名字','学号'), '第一行'),
                                       ([('甲','1')], ('姓名','姓名'), '第一行'),
                                       ([(f'学生{i}',str(i)) for i in range(501)], ('姓名','学号'), '最多导入500')]:
            self.assertIn(expected, self.roster_preview(rows, headers=headers).text)
        for filename in ['bad.xlsx','old.xls']:
            page = self.post(self.a, '/students/import', {'file':(io.BytesIO(b'not an excel file'), filename)})
            self.assertEqual(page.status_code, 200)
            self.assertIsNone(self.preview_token(page))
        page = self.post(self.a, '/students/import', {'file':(io.BytesIO(b'x'*(1024*1024)), 'large.xlsx')})
        self.assertEqual(page.status_code, 413)
        self.assertIn('1MB', page.text)
        from zipfile import ZipFile, ZIP_DEFLATED
        oversized = io.BytesIO()
        with ZipFile(oversized, 'w', compression=ZIP_DEFLATED) as archive:
            archive.writestr('large.xml', b'x'*(8*1024*1024+1))
        oversized.seek(0)
        page = self.post(self.a, '/students/import', {'file':(oversized, 'expanded.xlsx')})
        self.assertIn('内容过大', page.text)

    def test_import_numeric_codes_preserve_zero_format_and_optional_codes(self):
        from student_import import parse_roster
        wb = Workbook()
        wb.active.append(['学号','姓名'])
        wb.active.append([12,'带前导零'])
        wb.active['A2'].number_format = '00000'
        buf = io.BytesIO(); wb.save(buf)
        self.assertEqual(parse_roster(buf.getvalue())[0]['student_code'], '00012')
        page = self.roster_preview([('只有姓名',)], headers=('姓名',))
        self.assertTrue(self.preview_token(page))

    def test_summary_and_export_exclude_other_classes_but_keep_scored_history(self):
        self.save(score=88)
        self.post(self.a, '/create_class', {'name':'二班'})
        cid = self.sql("SELECT id FROM classes WHERE name='二班'")[0][0]
        self.post(self.a, '/add', {'name':'其他班学生','class_id':cid})
        context = self.summary()
        self.assertEqual({s['id'] for s in context['students']}, {self.sa,self.s2})
        self.assertEqual(context['statistics'][0]['expected'], 2)
        self.post(self.a, f'/edit_student/{self.sa}', {'name':'学生甲','class_id':cid})
        self.assertEqual({s['id'] for s in self.summary()['students']}, {self.sa,self.s2})
        exported = load_workbook(io.BytesIO(self.a.get('/export_excel').data))
        names = {r[0] for r in list(exported.active.values)[1:]}
        self.assertEqual(names, {'学生甲','学生乙'})

    def grid_data(self, client=None, exam=None, page=1):
        response = (client or self.a).get(f'/exam/{exam or self.ea}/scores?page={page}')
        self.assertEqual(response.status_code, 200)
        data = {name:html.unescape(value) for name,value in re.findall(r'name="(cell_\d+_\d+)" value="([^"]*)"', response.text)}
        token = re.search(r'name="snapshot" value="([^"]+)"', response.text)
        if token:
            data['snapshot'] = token[1]
        return data

    def test_class_scores_save_zero_absent_edit_clear_and_noop(self):
        data = self.grid_data()
        math_id = self.sql("SELECT id FROM exam_subjects WHERE exam_id=? AND subject='数学'", (self.ea,))[0][0]
        data[f'cell_{self.sa}_{self.ca}'] = '0'
        data[f'cell_{self.sa}_{math_id}'] = '缺考'
        data[f'cell_{self.s2}_{self.ca}'] = '99.5'
        self.assertFalse(self.sql('SELECT 1 FROM scores'))
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', data).status_code, 302)
        rows = self.sql('SELECT * FROM scores ORDER BY id')
        self.assertEqual(len(rows), 3)
        self.assertTrue(any(r['score'] == 0 and r['status'] == 'recorded' for r in rows))
        self.assertTrue(any(r['score'] is None and r['status'] == 'absent' for r in rows))
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', self.grid_data()).status_code, 302)
        self.assertEqual(len(self.sql('SELECT id FROM scores')), 3)
        data = self.grid_data()
        data[f'cell_{self.sa}_{self.ca}'] = ''
        data[f'cell_{self.s2}_{self.ca}'] = '100'
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', data).status_code, 302)
        self.assertFalse(self.sql('SELECT 1 FROM scores WHERE student_id=? AND subject=?', (self.sa,'语文')))
        self.assertEqual(self.sql('SELECT score FROM scores WHERE student_id=? AND subject=?', (self.s2,'语文'))[0][0], 100)

    def test_class_scores_invalid_cell_rolls_back_entire_page(self):
        for value in ['nan','inf','-1','151','abc','1e309']:
            with self.subTest(value=value):
                data = self.grid_data()
                data[f'cell_{self.sa}_{self.ca}'] = '90'
                data[f'cell_{self.s2}_{self.ca}'] = value
                page = self.post(self.a, f'/exam/{self.ea}/scores', data)
                self.assertEqual(page.status_code, 400)
                self.assertIn('没有保存任何修改', page.text)
                self.assertIn('value="90"', page.text)
                self.assertFalse(self.sql('SELECT id FROM scores'))

    def test_class_scores_conflict_preserves_newer_scores_and_entered_values(self):
        data = self.grid_data()
        data[f'cell_{self.sa}_{self.ca}'] = '91'
        self.save(score=88)
        response = self.post(self.a, f'/exam/{self.ea}/scores', data)
        self.assertEqual(response.status_code, 409)
        self.assertIn('value="91"', response.text)
        self.assertNotIn('name="snapshot"', response.text)
        self.assertEqual(self.sql('SELECT score FROM scores')[0][0], 88)
        for change in ['roster','config']:
            data = self.grid_data()
            data[f'cell_{self.sa}_{self.ca}'] = '92'
            if change == 'roster':
                self.post(self.a, '/add', {'name':'新加入学生'})
            else:
                self.sql('UPDATE exam_subjects SET max_score=200 WHERE id=?', (self.ca,))
            self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', data).status_code, 409)
            self.assertEqual(self.sql('SELECT score FROM scores')[0][0], 88)

    def test_class_scores_scope_csrf_tamper_and_expiry(self):
        self.assertEqual(self.b.get(f'/exam/{self.ea}/scores').status_code, 404)
        self.assertEqual(self.client().get(f'/exam/{self.ea}/scores').status_code, 302)
        data = self.grid_data()
        self.assertEqual(self.a.post(f'/exam/{self.ea}/scores', data=data).status_code, 400)
        self.assertEqual(self.post(self.b, f'/exam/{self.ea}/scores', data).status_code, 404)
        self.assertEqual(self.post(self.admin, f'/exam/{self.ea}/scores', data).status_code, 403)
        tampered = {**data,'snapshot':data['snapshot']+'bad'}
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', tampered).status_code, 400)
        forged = {**data,f'cell_{self.sb}_{self.ca}':'99'}
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', forged).status_code, 400)
        missing = dict(data); missing.pop(f'cell_{self.sa}_{self.ca}')
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', missing).status_code, 400)
        from itsdangerous import URLSafeTimedSerializer
        signer = URLSafeTimedSerializer(site.app.secret_key, salt='class-scores-v1')
        payload = signer.loads(data['snapshot'])
        with patch('itsdangerous.timed.time.time', return_value=time.time()-3601):
            expired = signer.dumps(payload)
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores', {**data,'snapshot':expired}).status_code, 400)
        self.assertFalse(self.sql('SELECT id FROM scores'))

    def test_class_scores_pagination_only_updates_requested_page(self):
        owner,cid = self.sql('SELECT teacher_id,class_id FROM students WHERE id=?', (self.sa,))[0]
        with site.database(write=True) as conn:
            conn.executemany('INSERT INTO students(name,teacher_id,class_id) VALUES(?,?,?)',
                             [(f'分页学生{i}', owner, cid) for i in range(31)])
        first, second = self.grid_data(), self.grid_data(page=2)
        second_keys = [k for k in second if k.startswith('cell_')]
        self.assertEqual(len(second_keys), 3*6)
        key = second_keys[0]
        second[key] = '88'
        self.assertEqual(self.post(self.a, f'/exam/{self.ea}/scores?page=2', second).status_code, 302)
        self.assertEqual(len(self.sql('SELECT id FROM scores')), 1)
        self.assertNotIn(key, first)
        for page, expected in [('0',400), ('bad',400), ('3',404)]:
            self.assertEqual(self.a.get(f'/exam/{self.ea}/scores?page={page}').status_code, expected)
        self.post(self.a, '/create_class', {'name':'转班分页'})
        target = self.sql("SELECT id FROM classes WHERE name='转班分页'")[0][0]
        tail_ids = [int(k.split('_')[1]) for k in second_keys if k.endswith('_'+str(self.ca))]
        for sid in tail_ids:
            self.save(sid=sid)
            self.sql('UPDATE students SET class_id=? WHERE id=?', (target,sid))
        data = self.grid_data(page=2)
        for field in data:
            if field.startswith('cell_'):
                data[field] = ''
        response = self.post(self.a, f'/exam/{self.ea}/scores?page=2', data)
        self.assertEqual(response.status_code, 302)
        self.assertIn('page=1', response.location)
        self.assertEqual(self.a.get(response.location).status_code, 200)

    def test_class_scores_excludes_unrelated_class_and_allows_historical_student(self):
        self.save(score=88)
        self.post(self.a, '/create_class', {'name':'录分二班'})
        cid = self.sql("SELECT id FROM classes WHERE name='录分二班'")[0][0]
        self.post(self.a, '/add', {'name':'无关班级学生','class_id':cid})
        unrelated = self.sql("SELECT id FROM students WHERE name='无关班级学生'")[0][0]
        page = self.a.get(f'/exam/{self.ea}/scores')
        self.assertNotIn('无关班级学生', page.text)
        self.assertEqual(self.save(sid=unrelated).status_code, 400)
        self.assertEqual(self.a.get(f'/score/{unrelated}?exam_id={self.ea}').status_code, 400)
        self.assertNotIn('期中考试', self.a.get(f'/score/{unrelated}').text)
        self.post(self.a, f'/edit_student/{self.sa}', {'name':'学生甲','class_id':cid})
        self.assertIn('学生甲', self.a.get(f'/exam/{self.ea}/scores').text)
        self.assertEqual(self.save(score=89).status_code, 302)
        score_id = self.sql('SELECT id FROM scores WHERE student_id=?', (self.sa,))[0][0]
        response = self.post(self.a, f'/delete_score/{score_id}', follow=True)
        self.assertEqual(response.status_code, 200)

    def test_class_scores_admin_can_save_and_empty_states(self):
        data = self.grid_data(client=self.admin)
        data[f'cell_{self.sa}_{self.ca}'] = '99'
        self.assertEqual(self.post(self.admin, f'/exam/{self.ea}/scores', data).status_code, 302)
        self.sql('DELETE FROM exam_subjects WHERE exam_id=?', (self.ea,))
        self.assertIn('尚未配置科目', self.a.get(f'/exam/{self.ea}/scores').text)
        self.sql('DELETE FROM scores WHERE exam_id=?', (self.ea,))
        self.post(self.a, f'/delete/{self.sa}'); self.post(self.a, f'/delete/{self.s2}')
        self.assertIn('暂无学生', self.a.get(f'/exam/{self.ea}/scores').text)

    def test_student_account_sees_only_own_scores(self):
        self.save(score=91)
        response=self.post(self.a,f'/student/{self.sa}/account',{'email':'student@example.invalid','password':'StudentPass123!'})
        self.assertEqual(response.status_code,302)
        student=self.client('student_'+str(self.sa),'StudentPass123!')
        self.assertEqual(student.get('/').location,'/change_password')
        self.post(student,'/change_password',{'old_password':'StudentPass123!','new_password':'StudentChanged123!','confirm_password':'StudentChanged123!'})
        self.assertEqual(student.get('/').location,f'/student/{self.sa}')
        own=student.get(f'/student/{self.sa}')
        self.assertEqual(own.status_code,200)
        self.assertIn('91',own.text)
        self.assertNotIn('继续录成绩',own.text)
        self.assertEqual(student.get(f'/student/{self.s2}').status_code,404)
        for path in ['/classes','/exams','/summary','/export_excel','/students/import','/students/import-template',f'/exam/{self.ea}/scores',f'/score/{self.sa}',f'/edit_student/{self.sa}']:
            self.assertEqual(student.get(path).status_code,403,path)

    def test_student_account_requires_unique_email_and_single_binding(self):
        self.post(self.a,f'/student/{self.sa}/account',{'email':'teacher_a@example.invalid','password':'StudentPass123!'})
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE role='student'"))
        self.post(self.a,f'/student/{self.sa}/account',{'email':'student@example.invalid','password':'StudentPass123!'})
        self.post(self.a,f'/student/{self.sa}/account',{'email':'second@example.invalid','password':'StudentPass123!'})
        self.assertEqual(len(self.sql('SELECT id FROM users WHERE student_id=?',(self.sa,))),1)
        self.post(self.a,f'/delete/{self.sa}')
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE email='student@example.invalid'"))

    def make_multi_parent(self):
        self.post(self.a, f'/student/{self.sa}/parent-account',
                  {'email':'multi-parent@example.invalid','password':'ParentPass123!'})
        uid = self.sql("SELECT id FROM users WHERE email='multi-parent@example.invalid'")[0][0]
        client = self.client('parent_'+str(self.sa), 'ParentPass123!')
        self.post(client, '/change_password', {'old_password':'ParentPass123!',
                  'new_password':'ChangedParent123!','confirm_password':'ChangedParent123!'})
        return uid, client

    def test_parent_multi_links_children_across_teachers_without_resetting_password(self):
        uid, parent = self.make_multi_parent()
        original = tuple(self.sql('SELECT password,session_version,must_change_password FROM users WHERE id=?', (uid,))[0])
        self.post(self.b, f'/student/{self.sb}/parent-account', {'email':'MULTI-PARENT@example.invalid'})
        self.assertEqual(tuple(self.sql('SELECT password,session_version,must_change_password FROM users WHERE id=?', (uid,))[0]), original)
        self.assertEqual(len(self.sql('SELECT * FROM parent_students WHERE parent_id=?', (uid,))), 2)
        home = parent.get('/')
        self.assertEqual(home.status_code, 200)
        self.assertIn('我的孩子', home.text)
        self.assertIn('其他老师学生', home.text)
        self.assertNotIn('学生乙', home.text)
        for sid in [self.sa,self.sb]:
            page = parent.get(f'/student/{sid}')
            self.assertEqual(page.status_code, 200)
            self.assertIn('切换孩子', page.text)
        self.assertEqual(parent.get(f'/student/{self.s2}').status_code, 404)
        self.assertIn('multi-parent@example.invalid', self.b.get(f'/edit_student/{self.sb}').text)
        self.assertNotIn('学生甲', self.b.get(f'/edit_student/{self.sb}').text)

    def test_parent_multi_delete_one_child_keeps_account_and_remaining_access(self):
        uid, parent = self.make_multi_parent()
        self.post(self.a, f'/student/{self.s2}/parent-account', {'email':'multi-parent@example.invalid'})
        self.sql('INSERT INTO reset_tokens(email,token,expires_at) VALUES(?,?,?)',
                 ('multi-parent@example.invalid','preserved-token',time.time()+300))
        self.post(self.a, f'/delete/{self.sa}')
        self.assertTrue(self.sql('SELECT 1 FROM users WHERE id=?', (uid,)))
        self.assertTrue(self.sql("SELECT 1 FROM reset_tokens WHERE email='multi-parent@example.invalid'"))
        self.assertEqual(parent.get('/').location, f'/student/{self.s2}')
        self.assertEqual(parent.get(f'/student/{self.sa}').status_code, 404)
        self.assertEqual(parent.get(f'/student/{self.s2}').status_code, 200)
        self.post(self.a, f'/delete/{self.s2}')
        self.assertFalse(self.sql('SELECT 1 FROM users WHERE id=?', (uid,)))
        self.assertFalse(self.sql('SELECT 1 FROM parent_students WHERE parent_id=?', (uid,)))
        self.assertFalse(self.sql("SELECT 1 FROM reset_tokens WHERE email='multi-parent@example.invalid'"))
        self.assertEqual(parent.get('/').location, '/login')

    def test_parent_multi_unlink_is_scoped_immediate_and_relinkable(self):
        uid, parent = self.make_multi_parent()
        self.post(self.a, f'/student/{self.s2}/parent-account', {'email':'multi-parent@example.invalid'})
        route = f'/student/{self.sa}/parent-account/{uid}/unlink'
        self.assertEqual(self.post(self.b, route).status_code, 404)
        self.assertEqual(self.post(parent, route).status_code, 403)
        self.assertEqual(self.a.post(route).status_code, 400)
        self.assertEqual(self.post(self.a, route).status_code, 302)
        self.assertEqual(parent.get(f'/student/{self.sa}').status_code, 404)
        self.assertEqual(parent.get(f'/student/{self.s2}').status_code, 200)
        self.assertEqual(self.post(self.a, route).status_code, 404)
        self.post(self.a, f'/student/{self.s2}/parent-account/{uid}/unlink')
        self.assertEqual(parent.get('/').status_code, 200)
        self.assertIn('目前没有关联孩子', parent.get('/').text)
        self.assertIn('退出', parent.get('/').text)
        self.assertTrue(self.sql('SELECT 1 FROM users WHERE id=?', (uid,)))
        self.post(self.a, f'/student/{self.sa}/parent-account', {'email':'multi-parent@example.invalid'})
        self.assertEqual(parent.get('/').location, f'/student/{self.sa}')
        self.assertEqual(self.sql('SELECT student_id FROM users WHERE id=?', (uid,))[0][0], self.sa)

    def test_parent_multi_duplicate_disabled_and_other_role_emails_rejected(self):
        uid, _ = self.make_multi_parent()
        self.post(self.a, f'/student/{self.sa}/parent-account', {'email':'multi-parent@example.invalid'})
        self.assertEqual(len(self.sql('SELECT 1 FROM parent_students WHERE parent_id=?', (uid,))), 1)
        self.post(self.admin, f'/delete_user/{uid}')
        self.post(self.a, f'/student/{self.s2}/parent-account', {'email':'multi-parent@example.invalid'})
        self.assertEqual(len(self.sql('SELECT 1 FROM parent_students WHERE parent_id=?', (uid,))), 1)
        self.post(self.a, f'/student/{self.s2}/parent-account', {'email':'teacher_a@example.invalid'})
        self.assertFalse(self.sql('SELECT 1 FROM parent_students WHERE student_id=?', (self.s2,)))
        self.post(self.a, f'/student/{self.s2}/parent-account', {'email':'new-parent@example.invalid'})
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE email='new-parent@example.invalid'"))

    def test_parent_accounts_see_only_bound_student(self):
        self.save(score=88)
        self.post(self.a, f'/student/{self.sa}/parent-account',
                  {'email':'parent1@example.invalid','password':'ParentPass123!'})
        self.post(self.a, f'/student/{self.sa}/parent-account',
                  {'email':'parent2@example.invalid','password':'ParentPass123!'})
        self.assertEqual(len(self.sql("SELECT id FROM users WHERE role='parent' AND student_id=?", (self.sa,))), 2)
        parent = self.client('parent_'+str(self.sa), 'ParentPass123!')
        self.post(parent, '/change_password', {'old_password':'ParentPass123!',
                  'new_password':'ParentChanged123!','confirm_password':'ParentChanged123!'})
        self.assertEqual(parent.get('/').location, f'/student/{self.sa}')
        own = parent.get(f'/student/{self.sa}')
        self.assertEqual(own.status_code, 200)
        self.assertIn('88', own.text)
        self.assertNotIn('继续录成绩', own.text)
        self.assertEqual(parent.get(f'/student/{self.s2}').status_code, 404)
        for path in ['/classes','/exams','/summary','/export_excel','/students/import','/students/import-template',f'/exam/{self.ea}/scores',f'/score/{self.sa}',f'/edit_student/{self.sa}']:
            self.assertEqual(parent.get(path).status_code, 403, path)

    def test_family_chart_excludes_unrelated_class_exams_and_keeps_history(self):
        self.save(score=88)
        self.post(self.a, '/create_class', {'name':'其他班级'})
        other_class = self.sql("SELECT id FROM classes WHERE name='其他班级'")[0][0]
        self.post(self.a, '/create_exam', {'name':'其他班级保密考试','class_id':other_class})
        self.post(self.a, '/create_class', {'name':'转入班级'})
        new_class = self.sql("SELECT id FROM classes WHERE name='转入班级'")[0][0]
        self.post(self.a, '/create_exam', {'name':'转入班级待考','class_id':new_class})
        self.post(self.a, f'/edit_student/{self.sa}',
                  {'name':'学生甲','class_id':new_class})
        for role, endpoint in [('student', 'account'), ('parent', 'parent-account')]:
            with self.subTest(role=role):
                self.post(self.a, f'/student/{self.sa}/{endpoint}',
                          {'email':role+'@example.invalid','password':'FamilyPass123!'})
                family = self.client(role+'_'+str(self.sa), 'FamilyPass123!')
                self.post(family, '/change_password', {'old_password':'FamilyPass123!',
                          'new_password':'ChangedFamily123!','confirm_password':'ChangedFamily123!'})
                page = family.get(f'/student/{self.sa}')
                self.assertEqual(page.status_code, 200)
                exams = json.loads(re.search(r'const exams = (.*);', page.text)[1])
                self.assertEqual({e['name'] for e in exams}, {'期中考试','转入班级待考'})
                self.assertIn('88', page.text)

    def test_family_write_guards_and_disabled_sessions(self):
        self.save(score=88)
        score_id = self.sql('SELECT id FROM scores')[0][0]
        for role, endpoint in [('student', 'account'), ('parent', 'parent-account')]:
            with self.subTest(role=role):
                self.post(self.a, f'/student/{self.sa}/{endpoint}',
                          {'email':role+'@example.invalid','password':'FamilyPass123!'})
                family = self.client(role+'_'+str(self.sa), 'FamilyPass123!')
                self.assertEqual(family.get(f'/student/{self.sa}').location, '/change_password')
                self.post(family, '/change_password', {'old_password':'FamilyPass123!',
                          'new_password':'ChangedFamily123!','confirm_password':'ChangedFamily123!'})
                for path, data in [('/add', {'name':'非法新增'}),
                                   (f'/delete/{self.sa}', {}),
                                   (f'/edit_student/{self.sa}', {'name':'非法修改'}),
                                   (f'/student/{self.sa}/parent-account', {}),
                                   (f'/student/{self.sa}/account', {}),
                                   (f'/score/{self.sa}', {'exam_id':self.ea,'subject':'语文','score':99}),
                                   (f'/edit_score/{score_id}', {'score':99}),
                                   (f'/delete_score/{score_id}', {}),
                                   ('/students/import', {'action':'confirm','preview_token':'forged'}),
                                   (f'/delete_exam/{self.ea}', {})]:
                    self.assertEqual(self.post(family, path, data).status_code, 403, path)
                self.assertEqual(self.sql('SELECT score FROM scores WHERE id=?', (score_id,))[0][0], 88)
                self.assertEqual(self.sql('SELECT name FROM students WHERE id=?', (self.sa,))[0][0], '学生甲')
                uid = self.sql('SELECT id FROM users WHERE username=?', (role+'_'+str(self.sa),))[0][0]
                self.post(self.admin, f'/delete_user/{uid}')
                self.assertEqual(family.get(f'/student/{self.sa}').location, '/login')
                self.post(self.admin, f'/enable_user/{uid}')
                self.assertEqual(family.get(f'/student/{self.sa}').location, '/login')
                family = self.client(role+'_'+str(self.sa), 'ChangedFamily123!')
                self.assertEqual(family.get(f'/student/{self.sa}').status_code, 200)
                self.post(family, '/logout')
                self.assertEqual(family.get(f'/student/{self.sa}').location, '/login')

    def test_parent_account_validates_email_and_is_deleted_with_student(self):
        self.post(self.a, f'/student/{self.sa}/parent-account',
                  {'email':'teacher_a@example.invalid','password':'ParentPass123!'})
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE role='parent'"))
        self.post(self.a, f'/student/{self.sa}/parent-account',
                  {'email':'parent@example.invalid','password':'ParentPass123!'})
        self.post(self.a, f'/student/{self.sa}/account',
                  {'email':'student-after-parent@example.invalid','password':'StudentPass123!'})
        self.assertTrue(self.sql("SELECT 1 FROM users WHERE role='student' AND student_id=?", (self.sa,)))
        self.post(self.a, f'/delete/{self.sa}')
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE email='parent@example.invalid'"))

    def test_admin_created_user_forced_to_change_password(self):
        self.post(self.admin,'/add_user',{'username':'new_teacher','email':'new@example.invalid','password':'InitialPass123!'})
        c=self.client('new_teacher','InitialPass123!')
        self.assertEqual(c.get('/').location,'/change_password')
        self.assertEqual(c.get('/export_excel').location,'/change_password')
        self.post(c,'/change_password',{'old_password':'InitialPass123!','new_password':'ChangedPass123!','confirm_password':'ChangedPass123!'})
        self.assertEqual(c.get('/').status_code,200)

    def test_admin_created_user_requires_unique_email(self):
        self.post(self.admin,'/add_user',{'username':'new_teacher','password':'InitialPass123!'})
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE username='new_teacher'"))
        self.post(self.admin,'/add_user',{'username':'new_teacher','email':'teacher_a@example.invalid','password':'InitialPass123!'})
        self.assertFalse(self.sql("SELECT 1 FROM users WHERE username='new_teacher'"))

    def test_cascade_student_exam_and_template_idempotent(self):
        self.save();self.post(self.a,f'/exam/{self.ea}/apply_template')
        self.assertEqual(len(self.sql('SELECT * FROM exam_subjects WHERE exam_id=?',(self.ea,))),6)
        self.post(self.a,f'/delete/{self.sa}');self.assertFalse(self.sql('SELECT 1 FROM scores'))
        self.post(self.a,f'/delete_exam/{self.ea}');self.assertFalse(self.sql('SELECT 1 FROM exam_subjects WHERE exam_id=?',(self.ea,)))

    def test_db_rejects_nonfinite_direct_write(self):
        for val in [float('nan'),float('inf'),-1,None]:
            with self.assertRaises(sqlite3.IntegrityError):
                self.sql('INSERT INTO scores(student_id,exam_id,subject,score) VALUES(?,?,?,?)',(self.sa,self.ea,'语文',val))

    def test_concurrent_overwrite_no_duplicate(self):
        clients=[self.client('teacher_a') for _ in range(2)]
        with ThreadPoolExecutor(max_workers=2) as pool:
            statuses=list(pool.map(lambda pair:self.save(client=pair[0],score=pair[1]).status_code,zip(clients,[90,95])))
        self.assertEqual(statuses,[302,302])
        self.assertEqual(len(self.sql('SELECT * FROM scores')),1)


    def test_csrf_unicode_is_rejected_without_server_error(self):
        response=self.a.post('/add',data={'name':'错误请求','csrf_token':'中文令牌'})
        self.assertEqual(response.status_code,400)

    def test_reserved_admin_and_known_password_rejected(self):
        c=self.client()
        response=self.post(c,'/register',dict(username='ADMIN',password='TeacherPass123!',confirm='TeacherPass123!'))
        self.assertIn('系统保留',response.text)
        response=self.post(self.a,'/change_password',dict(old_password='TeacherPass123!',new_password='admin123',confirm_password='admin123'))
        self.assertIn('公开默认密码',response.text)

    def test_resend_invalidates_old_code(self):
        with patch.object(site.secrets,'randbelow',return_value=123456):
            first=self.reset_code()
        self.sql('UPDATE rate_limits SET started_at=started_at-61')
        with patch.object(site.secrets,'randbelow',return_value=234567):
            second=self.reset_code()
        self.assertNotEqual(first,second)
        self.assertEqual(self.reset(first).status_code,200)
        self.assertEqual(self.reset(second).status_code,302)


class MigrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.db=self.root/'legacy.db'
    def tearDown(self):self.tmp.cleanup()
    def make_legacy(self):
        conn=sqlite3.connect(self.db)
        conn.executescript('''CREATE TABLE users(id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE,password TEXT,email TEXT);
          CREATE TABLE students(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT,teacher_id INTEGER);
          CREATE TABLE exams(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT UNIQUE,teacher_id INTEGER);
          CREATE TABLE scores(id INTEGER PRIMARY KEY AUTOINCREMENT,student_id INTEGER,subject TEXT,score REAL,exam TEXT,exam_id INTEGER);
          CREATE TABLE exam_subjects(id INTEGER PRIMARY KEY AUTOINCREMENT,exam_id INTEGER,subject TEXT,max_score REAL,pass_score REAL);
          CREATE TABLE reset_tokens(id INTEGER PRIMARY KEY AUTOINCREMENT,email TEXT,token TEXT,expires_at REAL);
          INSERT INTO students VALUES(10,'保留学生',7);INSERT INTO exams VALUES(20,'期中',7);
          INSERT INTO scores VALUES(30,10,'语文',90,'期中',20);INSERT INTO exam_subjects VALUES(40,20,'语文',150,90);''')
        conn.execute('INSERT INTO users VALUES(7,?,?,?)',('admin',generate_password_hash('admin123'),'admin@example.invalid'))
        conn.commit();conn.close()
    def migrate(self):init_database(self.db,self.root/'backups',self.root/'instance','InitialPass123!')

    def test_migration_preserves_data_and_backups_idempotent(self):
        self.make_legacy()
        conn=sqlite3.connect(self.db);conn.execute("INSERT INTO exams(id,name,teacher_id) VALUES(100,'已删考试',7)");conn.execute('DELETE FROM exams WHERE id=100');conn.commit();conn.close()
        self.migrate();self.migrate()
        conn=sqlite3.connect(self.db)
        self.assertEqual(conn.execute('SELECT id,score,exam_id FROM scores').fetchone(),(30,90,20))
        self.assertEqual(conn.execute('SELECT must_change_password FROM users').fetchone()[0],1)
        self.assertEqual(conn.execute("SELECT role FROM users WHERE username='admin'").fetchone()[0],'admin')
        cursor=conn.execute("INSERT INTO exams(name,teacher_id) VALUES('期中',8)");self.assertGreater(cursor.lastrowid,100);conn.commit();conn.close()
        backups=list((self.root/'backups').glob('*.db'));self.assertEqual(len(backups),1)
        original=sqlite3.connect(backups[0]);self.assertEqual(original.execute('SELECT score FROM scores').fetchone()[0],90);original.close()

    def test_parent_multi_v5_migration_preserves_accounts_tokens_and_business_rows(self):
        self.make_legacy(); self.migrate()
        conn = sqlite3.connect(self.db)
        conn.execute('DROP TABLE parent_students')
        conn.execute('UPDATE schema_meta SET version=5')
        conn.execute('INSERT INTO users(username,password,email,role,student_id,session_version,must_change_password) VALUES(?,?,?,?,?,?,?)',
                     ('legacy_parent',generate_password_hash('ExistingParent123!'),'legacy-parent@example.invalid','parent',10,9,0))
        uid = conn.execute("SELECT id FROM users WHERE username='legacy_parent'").fetchone()[0]
        conn.execute('INSERT INTO reset_tokens(email,token,expires_at,delivered) VALUES(?,?,?,1)',
                     ('legacy-parent@example.invalid','unchanged',time.time()+300))
        tables = ['users','students','classes','exams','exam_subjects','scores','reset_tokens']
        before = {table:conn.execute(f'SELECT * FROM {table} ORDER BY id').fetchall() for table in tables}
        conn.commit(); conn.close()
        self.migrate(); self.migrate()
        conn = sqlite3.connect(self.db)
        after = {table:conn.execute(f'SELECT * FROM {table} ORDER BY id').fetchall() for table in tables}
        self.assertEqual(before, after)
        self.assertEqual(conn.execute('SELECT * FROM parent_students').fetchall(), [(uid,10)])
        self.assertEqual(conn.execute('SELECT version FROM schema_meta').fetchone()[0], 6)
        self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(), [])
        conn.close()
        backups = list((self.root/'backups').glob('*.db'))
        self.assertEqual(len(backups), 2)
        version_five_backups = 0
        for path in backups:
            saved = sqlite3.connect(path)
            try:
                if saved.execute("SELECT 1 FROM sqlite_master WHERE name='schema_meta'").fetchone():
                    self.assertEqual(saved.execute('SELECT version FROM schema_meta').fetchone()[0], 5)
                    self.assertFalse(saved.execute("SELECT 1 FROM sqlite_master WHERE name='parent_students'").fetchone())
                    version_five_backups += 1
            finally:
                saved.close()
        self.assertEqual(version_five_backups, 1)

    def test_legacy_missing_config_and_nan_render_without_crash(self):
        self.make_legacy()
        conn=sqlite3.connect(self.db);conn.execute('DELETE FROM exam_subjects');conn.execute('UPDATE scores SET score=NULL');conn.execute('UPDATE users SET password=?',(generate_password_hash('LegacyPass123!'),));conn.commit();conn.close()
        self.migrate();old=site.DB;site.DB=str(self.db)
        try:
            c=site.app.test_client()
            with site.database() as db:
                user=db.execute("SELECT id,session_version FROM users WHERE username='admin'").fetchone()
            with c.session_transaction() as s:
                s.update(user='admin',uid=user['id'],session_version=user['session_version'],permanent=True)
            for path in ['/student/10','/score/10','/edit_score/30','/summary','/export_excel']:
                self.assertEqual(c.get(path).status_code,200,path)
            self.assertIn('数据异常',c.get('/student/10').text)
        finally:site.DB=old

    def test_pre_exam_id_migration(self):
        conn=sqlite3.connect(self.db);conn.executescript('''CREATE TABLE users(id INTEGER PRIMARY KEY,username TEXT UNIQUE,password TEXT);
            CREATE TABLE students(id INTEGER PRIMARY KEY,name TEXT);
            CREATE TABLE scores(id INTEGER PRIMARY KEY,student_id INTEGER,subject TEXT,score REAL,exam TEXT);
            INSERT INTO students VALUES(1,'旧学生'); INSERT INTO scores VALUES(1,1,'语文',100,'旧考试');''')
        conn.execute('INSERT INTO users VALUES(1,?,?)',('admin',generate_password_hash('LegacyPass123!')));conn.commit();conn.close()
        self.migrate();conn=sqlite3.connect(self.db)
        self.assertIsNotNone(conn.execute('SELECT exam_id FROM scores').fetchone()[0]);self.assertEqual(conn.execute('SELECT teacher_id FROM students').fetchone()[0],1)
        self.assertEqual(conn.execute('SELECT max_score FROM exam_subjects').fetchone()[0],150);conn.close()

    def test_bootstrap_random_password_not_public_default(self):
        init_database(self.db,self.root/'backups',self.root/'instance')
        pw=(self.root/'instance/admin-initial-password').read_text().strip()
        self.assertGreaterEqual(len(pw),8);self.assertNotEqual(pw,'admin123')
        conn=sqlite3.connect(self.db);row=conn.execute('SELECT password,must_change_password FROM users').fetchone();conn.close()
        self.assertTrue(check_password_hash(row[0],pw));self.assertEqual(row[1],1)
        self.assertEqual((self.root/'instance/admin-initial-password').stat().st_mode & 0o777,0o600)

    def test_backup_restore_integrity(self):
        self.make_legacy();target=backup_database(self.db,self.root/'backups')
        conn=sqlite3.connect(target);self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0],'ok');self.assertEqual(conn.execute('SELECT COUNT(*) FROM students').fetchone()[0],1);conn.close()


class StartupTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()
    def launch(self, mode='test', key=None, config_text=''):
        env=dict(os.environ)
        env.pop('SECRET_KEY',None)
        config=self.root/'env';config.write_text(config_text)
        env.update(CLASS_MANAGER_DB=str(self.root/'site.db'),CLASS_MANAGER_INSTANCE=str(self.root/'instance'),
                   CLASS_MANAGER_BACKUPS=str(self.root/'backups'),CLASS_MANAGER_ENV_FILE=str(config),CLASS_MANAGER_ENV=mode)
        if key is not None:env['SECRET_KEY']=key
        return subprocess.run([sys.executable,'-c',"import app,hashlib,json;print(json.dumps(dict(secure=app.app.config['SESSION_COOKIE_SECURE'],key_hash=hashlib.sha256(app.app.secret_key.encode()).hexdigest())))"],cwd=ROOT,env=env,capture_output=True,text=True)
    def test_production_requires_secret_and_rejects_old_default(self):
        self.assertNotEqual(self.launch('production').returncode,0)
        self.assertNotEqual(self.launch('production','short').returncode,0)
        self.assertNotEqual(self.launch('production','student-grade-system-secret-2026').returncode,0)
    def test_config_file_loaded_before_app_and_secure_cookie(self):
        result=self.launch('production',config_text='SECRET_KEY='+('x'*64))
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertTrue(json.loads(result.stdout)['secure'])
    def test_development_key_persists_across_restarts(self):
        a=self.launch();b=self.launch()
        self.assertEqual(a.returncode,0,a.stderr);self.assertEqual(b.returncode,0,b.stderr)
        self.assertEqual(json.loads(a.stdout)['key_hash'],json.loads(b.stdout)['key_hash'])
        self.assertEqual((self.root/'instance/secret-key').stat().st_mode & 0o777,0o600)


if __name__=='__main__':unittest.main()
