"""班主任学生成绩系统：账号、考试、成绩与导出。"""
import os
import re
import io
import json
import base64
import time
import secrets
import hmac
import hashlib
import smtplib
from time import time as unix_time
from datetime import datetime, timezone, timedelta
from pathlib import Path
from functools import wraps
from urllib import request as urlreq
from email.mime.text import MIMEText
from email.utils import formataddr
from flask import (Flask, request, render_template, session, redirect,
                   send_file, url_for, abort, flash, g)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from database import transaction, init_database, default_subjects, finite_number

BASE_DIR = Path(__file__).resolve().parent
# Load before constructing the application, including SECRET_KEY.
_env_file = Path(os.environ.get('CLASS_MANAGER_ENV_FILE', '~/.class_manager_env')).expanduser()
if _env_file.exists():
    for line in _env_file.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith('#') and '=' in line:
            k, v = line.split('=', 1)
            os.environ.setdefault(k.strip(), v.strip())

DB = str(Path(os.environ.get('CLASS_MANAGER_DB', BASE_DIR / 'class.db')).resolve())
INSTANCE_DIR = Path(os.environ.get('CLASS_MANAGER_INSTANCE', Path(DB).parent / '.instance'))
BACKUP_DIR = Path(os.environ.get('CLASS_MANAGER_BACKUPS', Path(DB).parent.parent / '.class_manager_backups'))
PRODUCTION = os.environ.get('CLASS_MANAGER_ENV') == 'production'
secret = os.environ.get('SECRET_KEY')
if secret == 'student-grade-system-secret-2026':
    raise RuntimeError('请更换旧版公开 SECRET_KEY')
if PRODUCTION and (not secret or len(secret) < 32):
    raise RuntimeError('生产环境必须配置至少32字符的随机 SECRET_KEY')
if not secret:
    INSTANCE_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
    secret_file = INSTANCE_DIR / 'secret-key'
    # Atomic replace from a complete private file avoids partially read keys.
    if not secret_file.exists():
        candidate = INSTANCE_DIR / ('secret-' + secrets.token_hex(8))
        with open(os.open(candidate, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), 'w') as output:
            output.write(secrets.token_hex(32))
        try:
            os.link(candidate, secret_file)
        except FileExistsError:
            pass
        finally:
            candidate.unlink()
    secret = secret_file.read_text().strip()
    if len(secret) < 32:
        raise RuntimeError('本地会话密钥文件无效')

app = Flask(__name__)
app.config.update(SECRET_KEY=secret, SESSION_COOKIE_HTTPONLY=True,
                  SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=PRODUCTION,
                  PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
                  MAX_CONTENT_LENGTH=1024 * 1024)
init_database(DB, BACKUP_DIR, INSTANCE_DIR, os.environ.get('ADMIN_INITIAL_PASSWORD'))


def database(write=False):
    return transaction(DB, write=write)


@app.template_filter('num')
def num_filter(value):
    try:
        number = finite_number(value)
    except ValueError:
        return '—'
    return int(number) if number.is_integer() else number


def csrf_token():
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_urlsafe(32)
    return session['csrf_token']


app.jinja_env.globals['csrf_token'] = csrf_token


@app.before_request
def check_request():
    g.user = None
    if 'uid' in session:
        with database() as conn:
            user = conn.execute('SELECT * FROM users WHERE id=? AND active=1', (session['uid'],)).fetchone()
        if not user or session.get('session_version') != user['session_version']:
            session.clear()
        else:
            g.user = user
            session['user'] = user['username']
    if request.method in ('POST', 'PUT', 'PATCH', 'DELETE'):
        actual = request.form.get('csrf_token', '')
        expected = session.get('csrf_token', '')
        if not actual or not expected or not hmac.compare_digest(actual.encode(), expected.encode()):
            abort(400, description='页面已过期或请求无法验证，请刷新页面后重试。')
    if g.user and g.user['must_change_password'] and request.endpoint not in ('change_password', 'logout', 'static'):
        return redirect(url_for('change_password'))


@app.after_request
def response_headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(429)
def request_error(error):
    return render_template('error.html', code=error.code, message=error.description), error.code


@app.errorhandler(500)
def server_error(error):
    return render_template('error.html', code=500, message='暂时无法完成操作，请稍后重试。'), 500


def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not g.user:
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped


def admin_required(fn):
    @wraps(fn)
    @login_required
    def wrapped(*args, **kwargs):
        if g.user['username'] != 'admin':
            abort(403, description='只有管理员可以执行此操作。')
        return fn(*args, **kwargs)
    return wrapped


def is_admin():
    return g.user is not None and g.user['username'] == 'admin'


def get_my_student(conn, sid):
    if is_admin():
        return conn.execute('SELECT * FROM students WHERE id=?', (sid,)).fetchone()
    return conn.execute('SELECT * FROM students WHERE id=? AND teacher_id=?', (sid, g.user['id'])).fetchone()


def get_my_exam(conn, eid):
    if is_admin():
        return conn.execute('SELECT * FROM exams WHERE id=?', (eid,)).fetchone()
    return conn.execute('SELECT * FROM exams WHERE id=? AND teacher_id=?', (eid, g.user['id'])).fetchone()


def require_student(conn, sid):
    student = get_my_student(conn, sid)
    if student is None:
        abort(404, description='没有这个学生，或你没有访问权限。')
    return student


def require_exam(conn, eid):
    exam = get_my_exam(conn, eid)
    if exam is None:
        abort(404, description='没有这场考试，或你没有访问权限。')
    return exam


def own_exams(conn):
    if is_admin():
        return conn.execute('SELECT * FROM exams ORDER BY id').fetchall()
    return conn.execute('SELECT * FROM exams WHERE teacher_id=? ORDER BY id', (g.user['id'],)).fetchall()


def own_classes(conn, include_inactive=False):
    clause = '' if include_inactive else ' AND active=1'
    if is_admin():
        return conn.execute(f'SELECT * FROM classes WHERE 1=1{clause} ORDER BY teacher_id,id').fetchall()
    return conn.execute(f'SELECT * FROM classes WHERE teacher_id=?{clause} ORDER BY id', (g.user['id'],)).fetchall()


def get_my_class(conn, cid):
    if is_admin():
        return conn.execute('SELECT * FROM classes WHERE id=? AND active=1', (cid,)).fetchone()
    return conn.execute('SELECT * FROM classes WHERE id=? AND teacher_id=? AND active=1', (cid, g.user['id'])).fetchone()


def default_class(conn):
    row = conn.execute('SELECT * FROM classes WHERE teacher_id=? AND active=1 ORDER BY id LIMIT 1', (g.user['id'],)).fetchone()
    if row:
        return row
    cid = conn.execute('INSERT INTO classes(teacher_id,name,created_at) VALUES(?,?,?)', (g.user['id'], '默认班级', unix_time())).lastrowid
    return conn.execute('SELECT * FROM classes WHERE id=?', (cid,)).fetchone()


def requested_class(conn, value=None):
    try:
        cid = int(value) if value not in (None, '') else None
    except (TypeError, ValueError):
        abort(400, description='班级编号无效。')
    row = get_my_class(conn, cid) if cid is not None else default_class(conn)
    if row is None:
        abort(404, description='没有这个班级，或你没有访问权限。')
    return row


def consume_limits(limits):
    """Reserve persistent counters in one transaction, shared by web workers."""
    now = time.time()
    with database(write=True) as conn:
        conn.execute('DELETE FROM rate_limits WHERE started_at<?', (now - 86400,))
        rows = []
        for label, value, maximum, window in limits:
            key = hashlib.sha256((label + ':' + str(value)).encode()).hexdigest()
            row = conn.execute('SELECT * FROM rate_limits WHERE bucket=?', (key,)).fetchone()
            count = row['count'] if row and now < row['started_at'] + window else 0
            if count >= maximum:
                return False
            rows.append((key, count + 1, row['started_at'] if count else now))
        for key, count, started in rows:
            conn.execute('INSERT INTO rate_limits(bucket,count,started_at) VALUES(?,?,?) ON CONFLICT(bucket) DO UPDATE SET count=excluded.count,started_at=excluded.started_at', (key, count, started))
    return True


def valid_email(value):
    return bool(len(value) <= 254 and re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', value))


def password_error(value):
    if value == 'admin123':
        return '不能使用旧版公开默认密码，请设置新的密码'
    return '密码长度需为8至128位' if not 8 <= len(value) <= 128 else None


def send_email(to_addr, subject, body):
    """统一发信入口：优先 Google Apps Script(Gmail) → Brevo → Mailjet → SMTP"""
    if os.environ.get("GAPPS_URL"):
        return send_email_via_gapps(to_addr, subject, body)
    if os.environ.get("BREVO_API_KEY"):
        return send_email_via_brevo(to_addr, subject, body)
    if os.environ.get("MAILJET_KEY"):
        return send_email_via_mailjet(to_addr, subject, body)
    return send_email_via_smtp(to_addr, subject, body)


def send_email_via_gapps(to_addr, subject, body):
    """通过 Google Apps Script(Gmail) 发邮件（免费，每天100封）"""
    from urllib.parse import urlencode
    base = os.environ.get("GAPPS_URL")
    url = base + "?" + urlencode({"to": to_addr, "subject": subject, "body": body})
    with urlreq.urlopen(url, timeout=30) as resp:
        return resp.status


def send_email_via_brevo(to_addr, subject, body):
    """通过 Brevo HTTP API 发邮件（免费 300 封/天）"""
    key = os.environ.get("BREVO_API_KEY")
    sender = os.environ.get("BREVO_SENDER", "qq1039535644@gmail.com")
    payload = {
        "sender": {"email": sender, "name": "班主任成绩系统"},
        "to": [{"email": to_addr}],
        "subject": subject,
        "textContent": body
    }
    req = urlreq.Request(
        "https://api.brevo.com/v3/smtp/email",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json",
                 "api-key": key})
    with urlreq.urlopen(req, timeout=30) as resp:
        return resp.status


def send_email_via_mailjet(to_addr, subject, body):
    """通过 Mailjet HTTP API 发邮件（本地/公网都能用，公网在白名单内）"""
    key = os.environ.get("MAILJET_KEY")
    secret = os.environ.get("MAILJET_SECRET")
    sender = os.environ.get("MAILJET_SENDER")     # 已验证的发件邮箱
    if not (key and secret and sender):
        raise RuntimeError("未配置 Mailjet：需要 MAILJET_KEY / MAILJET_SECRET / MAILJET_SENDER")
    payload = {"Messages": [{
        "From": {"Email": sender, "Name": "班主任成绩系统"},
        "To": [{"Email": to_addr}],
        "Subject": subject, "TextPart": body}]}
    auth = base64.b64encode(f"{key}:{secret}".encode()).decode()
    req = urlreq.Request(
        "https://api.mailjet.com/v3.1/send",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json",
                 "Authorization": "Basic " + auth})
    with urlreq.urlopen(req, timeout=30) as resp:
        return resp.status


def send_email_via_smtp(to_addr, subject, body):
    """通过 SMTP 发邮件（支持 QQ/163/Gmail 等任意邮箱）。
    账号、授权码、服务器地址都从环境变量读取，不写死在代码里
    （防止传到 GitHub 泄露）。"""
    user = os.environ.get("SMTP_USER")      # 发件邮箱
    pwd = os.environ.get("SMTP_PASS")       # 授权码 / 应用专用密码
    host = os.environ.get("SMTP_HOST", "smtp.qq.com")   # SMTP 服务器
    port = int(os.environ.get("SMTP_PORT", "465"))      # 端口
    if not user or not pwd:
        raise RuntimeError("未配置发件邮箱：请先设置环境变量 SMTP_USER 和 SMTP_PASS")
    msg = MIMEText(body, "plain", "utf-8")
    msg["Subject"] = subject
    msg["From"] = formataddr(("班主任成绩系统", user))   # 发件人显示成名称
    msg["To"] = to_addr
    with smtplib.SMTP_SSL(host, port, timeout=30) as s:
        s.login(user, pwd)
        s.sendmail(user, [to_addr], msg.as_string())


@app.route('/login', methods=['GET', 'POST'])
def login():
    error = None
    info = '注册成功，请登录' if request.args.get('registered') else None
    if request.args.get('reset'):
        info = '密码已重置，请用新密码登录'
    if request.method == 'POST':
        identifier = request.form.get('email', request.form.get('username', '')).strip().lower()
        password = request.form.get('password', '')
        if not consume_limits([('login_ip', request.remote_addr, 60, 900), ('login_user', identifier.casefold(), 10, 900)]):
            return render_template('login.html', error='尝试次数过多，请15分钟后重试', info=None), 429
        with database() as conn:
            user = conn.execute('SELECT * FROM users WHERE lower(email)=? AND active=1', (identifier,)).fetchone()
            if not user:
                # Existing accounts without an email remain accessible until one is added.
                user = conn.execute('SELECT * FROM users WHERE username=? AND (email IS NULL OR email="") AND active=1', (identifier,)).fetchone()
        if len(password) <= 128 and user and check_password_hash(user['password'], password):
            session.clear()
            session.update(user=user['username'], uid=user['id'], session_version=user['session_version'])
            session.permanent = True
            if not user['email']:
                flash('该账号尚未登记邮箱。目前可继续用用户名登录；请联系管理员补充邮箱。')
            return redirect(url_for('change_password') if user['must_change_password'] else url_for('index'))
        error = '邮箱或密码错误，或账号已停用'
    return render_template('login.html', error=error, info=info)


@app.route('/logout', methods=['POST'])
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/register', methods=['GET', 'POST'])
def register():
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        email = request.form.get('email', '').strip().lower()
        if not consume_limits([('register_ip', request.remote_addr, 10, 3600)]):
            abort(429, description='注册过于频繁，请稍后重试。')
        with database(write=True) as conn:
            if username.casefold() == 'admin':
                error = '该用户名为系统保留账号'
            elif not 1 <= len(username) <= 50:
                error = '用户名长度需为1至50位'
            elif password_error(password):
                error = password_error(password)
            elif password != request.form.get('confirm', ''):
                error = '两次输入的密码不一致'
            elif not valid_email(email):
                error = '邮箱格式不正确'
            elif conn.execute('SELECT 1 FROM users WHERE username=?', (username,)).fetchone():
                error = '用户名已被占用，换一个'
            elif email and conn.execute('SELECT 1 FROM users WHERE lower(email)=?', (email,)).fetchone():
                error = '该邮箱已被注册'
            else:
                conn.execute('INSERT INTO users(username,password,email) VALUES(?,?,?)', (username, generate_password_hash(password), email))
                return redirect('/login?registered=1')
    return render_template('register.html', error=error)


@app.route('/forgot_password', methods=['GET', 'POST'])
def forgot_password():
    error = None
    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        if not valid_email(email):
            error = '邮箱格式不正确'
        elif not consume_limits([
            ('reset_cooldown', email, 1, 60), ('reset_email', email, 5, 3600),
            ('reset_ip', request.remote_addr, 15, 3600),
        ]):
            return render_template('forgot_password.html', error='发送过于频繁，请至少等待一分钟后再试', info=None), 429
        else:
            token = f'{secrets.randbelow(1000000):06d}'
            with database(write=True) as conn:
                conn.execute('DELETE FROM reset_tokens WHERE expires_at<? OR lower(email)=?', (time.time(), email))
                user = conn.execute('SELECT id FROM users WHERE lower(email)=? AND active=1', (email,)).fetchone()
                token_id = None
                if user:
                    token_id = conn.execute('INSERT INTO reset_tokens(email,token,expires_at,created_at) VALUES(?,?,?,?)', (email, generate_password_hash(token), time.time() + 600, time.time())).lastrowid
            if token_id is not None:
                try:
                    send_email(email, '成绩系统 · 密码重置验证码', f'您正在重置密码，验证码：{token}\n验证码10分钟内有效，最多尝试5次。若非本人操作，请忽略此邮件。')
                except Exception:
                    # Do not log provider URLs or bodies containing addresses, keys or codes.
                    app.logger.error('Password reset email could not be delivered')
                    with database(write=True) as conn:
                        conn.execute('DELETE FROM reset_tokens WHERE id=?', (token_id,))
                    error = '邮件暂时无法发送，请稍后重试或联系管理员'
                else:
                    with database(write=True) as conn:
                        conn.execute('UPDATE reset_tokens SET delivered=1 WHERE id=?', (token_id,))
            if not error:
                flash('如果该邮箱已注册且账号可用，验证码将发送到邮箱。请检查收件箱及垃圾箱。')
                return redirect(url_for('reset_password', email=email))
    return render_template('forgot_password.html', error=error, info=None)


@app.route('/reset_password', methods=['GET', 'POST'])
def reset_password():
    error = None
    email = request.values.get('email', '').strip().lower()
    if request.method == 'POST':
        if not consume_limits([('reset_verify_ip', request.remote_addr, 30, 900)]):
            abort(429, description='验证过于频繁，请稍后重试。')
        token = request.form.get('token', '').strip()
        new = request.form.get('new_password', '')
        with database(write=True) as conn:
            row = conn.execute('SELECT * FROM reset_tokens WHERE email=? ORDER BY id DESC LIMIT 1', (email,)).fetchone()
            if not row or not row['delivered'] or row['expires_at'] <= time.time() or row['attempts'] >= 5:
                error = '验证码无效或已过期，请重新获取'
            elif not re.fullmatch(r'\d{6}', token) or not check_password_hash(row['token'], token):
                conn.execute('UPDATE reset_tokens SET attempts=attempts+1 WHERE id=?', (row['id'],))
                error = '验证码不正确，连续错误5次需重新获取'
            elif password_error(new):
                error = password_error(new)
            elif new != request.form.get('confirm', ''):
                error = '两次输入的新密码不一致'
            else:
                # User and token changes are serialized so one code cannot be consumed twice.
                conn.execute('UPDATE users SET password=?,session_version=session_version+1,must_change_password=0 WHERE lower(email)=? AND active=1', (generate_password_hash(new), email))
                conn.execute('DELETE FROM reset_tokens WHERE email=?', (email,))
                session.clear()
                return redirect('/login?reset=1')
    return render_template('reset_password.html', error=error, email=email)


@app.route('/change_password', methods=['GET', 'POST'])
@login_required
def change_password():
    error = None
    info = '为保护账号，请先修改初始密码' if g.user['must_change_password'] else None
    if request.method == 'POST':
        new = request.form.get('new_password', '')
        old = request.form.get('old_password', '')
        with database(write=True) as conn:
            user = conn.execute('SELECT * FROM users WHERE id=?', (g.user['id'],)).fetchone()
            if len(old) > 128 or not check_password_hash(user['password'], old):
                error = '旧密码不正确'
            elif password_error(new):
                error = password_error(new)
            elif new == old:
                error = '新密码不能与旧密码相同'
            elif new != request.form.get('confirm_password', ''):
                error = '两次输入的新密码不一致'
            else:
                conn.execute('UPDATE users SET password=?,session_version=session_version+1,must_change_password=0 WHERE id=?', (generate_password_hash(new), user['id']))
                session['session_version'] = user['session_version'] + 1
                flash('密码修改成功，其他设备的登录已失效。')
                return redirect(url_for('index'))
    return render_template('change_password.html', error=error, info=info)


@app.route('/users')
@admin_required
def user_list():
    with database() as conn:
        users = conn.execute('SELECT id,username,email,active FROM users ORDER BY id').fetchall()
    return render_template('users.html', users=users, error=None)


@app.route('/add_user', methods=['POST'])
@admin_required
def add_user():
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '')
    email = request.form.get('email', '').strip().lower()
    with database(write=True) as conn:
        if not 1 <= len(username) <= 50 or password_error(password) or not valid_email(email):
            flash('用户名需为1至50位，邮箱格式正确，密码需为8至128位')
        elif username.casefold() == 'admin' or conn.execute('SELECT 1 FROM users WHERE username=?', (username,)).fetchone():
            flash('用户名已存在或为系统保留账号')
        elif conn.execute('SELECT 1 FROM users WHERE lower(email)=?', (email,)).fetchone():
            flash('该邮箱已被其他账号使用')
        else:
            conn.execute('INSERT INTO users(username,password,email,must_change_password) VALUES(?,?,?,1)', (username, generate_password_hash(password), email))
            flash('老师账号已创建，首次登录需修改密码。')
    return redirect(url_for('user_list'))


@app.route('/delete_user/<int:uid>', methods=['POST'])
@admin_required
def delete_user(uid):
    with database(write=True) as conn:
        conn.execute("UPDATE users SET active=0,session_version=session_version+1 WHERE id=? AND username!='admin'", (uid,))
        conn.execute('DELETE FROM reset_tokens WHERE lower(email) IN (SELECT lower(email) FROM users WHERE id=? AND active=0)', (uid,))
    flash('账号已停用，已有登录已失效，学生与考试数据已保留。')
    return redirect(url_for('user_list'))


@app.route('/enable_user/<int:uid>', methods=['POST'])
@admin_required
def enable_user(uid):
    with database(write=True) as conn:
        conn.execute("UPDATE users SET active=1,session_version=session_version+1 WHERE id=? AND username!='admin'", (uid,))
    flash('账号已恢复。')
    return redirect(url_for('user_list'))


@app.route('/')
@login_required
def index():
    with database() as conn:
        classes = own_classes(conn)
        cur_class = requested_class(conn, request.args.get('class_id'))
        students = conn.execute('SELECT * FROM students WHERE class_id=? ORDER BY id', (cur_class['id'],)).fetchall()
    return render_template('index.html', students=students, classes=classes, cur_class=cur_class)


@app.route('/add', methods=['POST'])
@login_required
def add_student():
    name = request.form.get('name', '').strip()
    student_code = request.form.get('student_code', '').strip()
    if not 1 <= len(name) <= 100:
        flash('学生姓名长度需为1至100位')
    elif len(student_code) > 50:
        flash('学号长度不能超过50位')
    else:
        with database(write=True) as conn:
            cur_class = requested_class(conn, request.form.get('class_id'))
            if conn.execute('SELECT 1 FROM students WHERE name=? AND class_id=?', (name, cur_class['id'])).fetchone():
                flash(f'已存在同名学生：{name}')
            elif student_code and conn.execute('SELECT 1 FROM students WHERE student_code=? AND class_id=?', (student_code, cur_class['id'])).fetchone():
                flash(f'该班级已有学号：{student_code}')
            else:
                conn.execute('INSERT INTO students(name,teacher_id,class_id,student_code) VALUES(?,?,?,?)', (name, cur_class['teacher_id'], cur_class['id'], student_code or None))
    return redirect(url_for('index', class_id=request.form.get('class_id') or ''))


@app.route('/classes')
@login_required
def class_list():
    with database() as conn:
        classes = own_classes(conn)
        counts = {c['id']: conn.execute('SELECT COUNT(*) FROM students WHERE class_id=?', (c['id'],)).fetchone()[0] for c in classes}
    return render_template('classes.html', classes=classes, counts=counts)


@app.route('/create_class', methods=['POST'])
@login_required
def create_class():
    name = request.form.get('name', '').strip()
    grade = request.form.get('grade', '').strip() or None
    school_year = request.form.get('school_year', '').strip() or None
    with database(write=True) as conn:
        if not 1 <= len(name) <= 80:
            flash('班级名称长度需为1至80位')
        elif conn.execute('SELECT 1 FROM classes WHERE teacher_id=? AND name=?', (g.user['id'], name)).fetchone():
            flash('你已经有同名班级')
        else:
            conn.execute('INSERT INTO classes(teacher_id,name,grade,school_year,created_at) VALUES(?,?,?,?,?)', (g.user['id'], name, grade, school_year, unix_time()))
            flash('班级已创建。')
    return redirect(url_for('class_list'))


@app.route('/select_class', methods=['POST'])
@login_required
def select_class():
    with database() as conn:
        cur_class = requested_class(conn, request.form.get('class_id'))
    return redirect(url_for('index', class_id=cur_class['id']))

def normalise_score(row):
    record = dict(row)
    record.update(valid=False, config_valid=False, fail=False, display='数据异常')
    try:
        maximum = finite_number(record.get('max_score'))
        passing = finite_number(record.get('pass_score'))
        if maximum <= 0 or not 0 <= passing <= maximum:
            return record
        record['config_valid'] = True
        if record.get('status') == 'absent':
            record.update(valid=True, display='缺考', score=None)
            return record
        score = finite_number(record.get('score'))
        if record.get('status') != 'recorded' or not 0 <= score <= maximum:
            return record
        record.update(valid=True, display=num_filter(score), score=score, fail=score < passing)
    except ValueError:
        pass
    return record


def get_student_scores(conn, sid):
    # Newest legacy duplicate wins; new writes update that row under a write lock.
    rows = conn.execute('''SELECT s.*,e.name exam_name,es.max_score,es.pass_score
        FROM scores s JOIN students st ON st.id=s.student_id
        JOIN exams e ON e.id=s.exam_id AND e.teacher_id=st.teacher_id
        LEFT JOIN exam_subjects es ON es.id=(SELECT MAX(cfg.id) FROM exam_subjects cfg WHERE cfg.exam_id=s.exam_id AND cfg.subject=s.subject)
        WHERE s.student_id=? AND s.id=(SELECT MAX(sc.id) FROM scores sc WHERE sc.student_id=s.student_id AND sc.exam_id=s.exam_id AND sc.subject=s.subject)
        ORDER BY s.exam_id,s.subject''', (sid,)).fetchall()
    return [normalise_score(row) for row in rows]


def require_pair(student, exam):
    if student['teacher_id'] != exam['teacher_id']:
        abort(400, description='学生与考试必须属于同一位老师。')


def parse_score(config):
    status = request.form.get('status', 'recorded')
    if status == 'absent':
        return None, status
    if status != 'recorded':
        raise ValueError('请选择正常成绩或缺考')
    maximum = finite_number(config['max_score'])
    passing = finite_number(config['pass_score'])
    if maximum <= 0 or not 0 <= passing <= maximum:
        raise ValueError('科目配置异常，请先修正满分和及格线')
    score = finite_number(request.form.get('score', '').strip())
    if not 0 <= score <= maximum:
        raise ValueError(f'请输入0至{num_filter(maximum)}分之间的成绩')
    return score, status


@app.route('/score/<int:sid>', methods=['GET', 'POST'])
@login_required
def score_page(sid):
    error = None
    with database(write=request.method == 'POST') as conn:
        student = require_student(conn, sid)
        exams = [e for e in own_exams(conn) if e['teacher_id'] == student['teacher_id']]
        eid = request.values.get('exam_id', type=int)
        if 'exam_id' in request.values and eid is None:
            abort(400, description='考试编号无效。')
        if eid is None and exams:
            eid = exams[0]['id']
        exam = require_exam(conn, eid) if eid is not None else None
        if exam:
            require_pair(student, exam)
        subs = conn.execute('SELECT * FROM exam_subjects WHERE exam_id=? ORDER BY id', (eid,)).fetchall()
        if request.method == 'POST':
            subject = request.form.get('subject', '').strip()
            config = conn.execute('SELECT * FROM exam_subjects WHERE exam_id=? AND subject=? ORDER BY id DESC LIMIT 1', (eid, subject)).fetchone()
            try:
                if not config:
                    raise ValueError('该考试没有配置此科目，请先添加科目')
                value, status = parse_score(config)
            except ValueError as exc:
                error = str(exc)
            else:
                existing = conn.execute('SELECT id FROM scores WHERE student_id=? AND exam_id=? AND subject=? ORDER BY id DESC', (sid, eid, subject)).fetchall()
                if existing:
                    conn.execute('UPDATE scores SET score=?,status=? WHERE id=?', (value, status, existing[0]['id']))
                    conn.execute('DELETE FROM scores WHERE student_id=? AND exam_id=? AND subject=? AND id!=?', (sid, eid, subject, existing[0]['id']))
                else:
                    conn.execute('INSERT INTO scores(student_id,exam_id,subject,score,status,exam) VALUES(?,?,?,?,?,?)', (sid, eid, subject, value, status, exam['name']))
                flash('成绩已保存。')
                return redirect(url_for('score_page', sid=sid, exam_id=eid))
        rows = get_student_scores(conn, sid)
    return render_template('score.html', stu=student, rows=rows, all_exams=exams, cur_exam=exam, cur_subs=subs, error=error)


@app.route('/student/<int:sid>')
@login_required
def student_page(sid):
    with database() as conn:
        student = require_student(conn, sid)
        rows = get_student_scores(conn, sid)
        exams = [dict(e) for e in own_exams(conn) if e['teacher_id'] == student['teacher_id']]
    chart = {}
    for row in rows:
        if row['valid'] and row['status'] == 'recorded':
            chart.setdefault(row['subject'], {})[str(row['exam_id'])] = round(row['score'] / row['max_score'] * 100, 2)
    return render_template('student.html', stu=student, rows=rows, all_exams=exams, chart=chart,
                           print_date=datetime.now(timezone(timedelta(hours=8))).strftime('%Y-%m-%d'))


@app.route('/help')
@login_required
def help_page():
    return render_template('help.html')


@app.route('/exams')
@login_required
def exam_list():
    with database() as conn:
        classes = own_classes(conn)
        cur_class = requested_class(conn, request.args.get('class_id'))
        exams = conn.execute('SELECT * FROM exams WHERE class_id=? ORDER BY id', (cur_class['id'],)).fetchall()
        counts = {e['id']: conn.execute('SELECT COUNT(*) FROM scores WHERE exam_id=?', (e['id'],)).fetchone()[0] for e in exams}
    return render_template('exams.html', exams=exams, counts=counts, classes=classes, cur_class=cur_class, error=None)


@app.route('/create_exam', methods=['POST'])
@login_required
def create_exam():
    name = request.form.get('name', '').strip()
    with database(write=True) as conn:
        cur_class = requested_class(conn, request.form.get('class_id'))
        if not 1 <= len(name) <= 100:
            flash('考试名称长度需为1至100位')
        elif conn.execute('SELECT 1 FROM exams WHERE class_id=? AND name=?', (cur_class['id'], name)).fetchone():
            flash('已存在同名考试')
        else:
            eid = conn.execute('INSERT INTO exams(name,teacher_id,class_id) VALUES(?,?,?)', (name, cur_class['teacher_id'], cur_class['id'])).lastrowid
            conn.executemany('INSERT INTO exam_subjects(exam_id,subject,max_score,pass_score) VALUES(?,?,?,?)', [(eid, subject, maximum, passing) for subject, maximum, passing in default_subjects()])
            return redirect(url_for('exam_config', exam_id=eid))
    return redirect(url_for('exam_list', class_id=request.form.get('class_id') or ''))


@app.route('/delete_exam/<int:exam_id>', methods=['POST'])
@login_required
def delete_exam(exam_id):
    with database(write=True) as conn:
        require_exam(conn, exam_id)
        conn.execute('DELETE FROM scores WHERE exam_id=?', (exam_id,))
        conn.execute('DELETE FROM exam_subjects WHERE exam_id=?', (exam_id,))
        conn.execute('DELETE FROM exams WHERE id=?', (exam_id,))
    return redirect(url_for('exam_list'))


@app.route('/exam/<int:exam_id>')
@login_required
def exam_config(exam_id):
    with database() as conn:
        exam = require_exam(conn, exam_id)
        subs = conn.execute('SELECT * FROM exam_subjects WHERE exam_id=? ORDER BY id', (exam_id,)).fetchall()
    return render_template('exam_config.html', exam=exam, subs=subs, error=None)


def parse_subject():
    maximum = finite_number(request.form.get('max_score', ''))
    raw = request.form.get('pass_score', '').strip()
    passing = finite_number(raw) if raw else maximum * .6
    if maximum <= 0 or not 0 <= passing <= maximum:
        raise ValueError('满分必须大于0，及格线需在0至满分之间')
    return maximum, passing


@app.route('/exam/<int:exam_id>/add_subject', methods=['POST'])
@login_required
def add_subject(exam_id):
    with database(write=True) as conn:
        require_exam(conn, exam_id)
        try:
            subject = request.form.get('subject', '').strip()
            maximum, passing = parse_subject()
            if not 1 <= len(subject) <= 50:
                raise ValueError('科目名称长度需为1至50位')
            if conn.execute('SELECT 1 FROM exam_subjects WHERE exam_id=? AND subject=?', (exam_id, subject)).fetchone():
                raise ValueError('该考试已有此科目')
            conn.execute('INSERT INTO exam_subjects(exam_id,subject,max_score,pass_score) VALUES(?,?,?,?)', (exam_id, subject, maximum, passing))
        except ValueError as exc:
            flash(str(exc))
    return redirect(url_for('exam_config', exam_id=exam_id))


@app.route('/exam/<int:exam_id>/update_subject/<int:cfg_id>', methods=['POST'])
@login_required
def update_subject(exam_id, cfg_id):
    with database(write=True) as conn:
        require_exam(conn, exam_id)
        config = conn.execute('SELECT * FROM exam_subjects WHERE id=? AND exam_id=?', (cfg_id, exam_id)).fetchone()
        if not config:
            abort(404, description='该科目不属于这场考试。')
        try:
            maximum, passing = parse_subject()
            scores = conn.execute("SELECT score FROM scores WHERE exam_id=? AND subject=? AND status='recorded'", (exam_id, config['subject'])).fetchall()
            for row in scores:
                try:
                    existing = finite_number(row['score'])
                except ValueError:
                    continue
                if existing > maximum:
                    raise ValueError(f'已有成绩为{num_filter(existing)}分，不能将满分下调到{num_filter(maximum)}分，请先修正相关成绩')
            conn.execute('UPDATE exam_subjects SET max_score=?,pass_score=? WHERE id=? AND exam_id=?', (maximum, passing, cfg_id, exam_id))
        except ValueError as exc:
            flash(str(exc))
    return redirect(url_for('exam_config', exam_id=exam_id))


@app.route('/exam/<int:exam_id>/delete_subject/<int:cfg_id>', methods=['POST'])
@login_required
def delete_subject(exam_id, cfg_id):
    with database(write=True) as conn:
        require_exam(conn, exam_id)
        config = conn.execute('SELECT * FROM exam_subjects WHERE id=? AND exam_id=?', (cfg_id, exam_id)).fetchone()
        if not config:
            abort(404, description='该科目不属于这场考试。')
        # Derive subject from the verified row, never from the request body.
        conn.execute('DELETE FROM scores WHERE exam_id=? AND subject=?', (exam_id, config['subject']))
        conn.execute('DELETE FROM exam_subjects WHERE exam_id=? AND subject=?', (exam_id, config['subject']))
    return redirect(url_for('exam_config', exam_id=exam_id))


@app.route('/exam/<int:exam_id>/apply_template', methods=['POST'])
@login_required
def apply_template(exam_id):
    with database(write=True) as conn:
        require_exam(conn, exam_id)
        for subject, maximum, passing in default_subjects():
            if not conn.execute('SELECT 1 FROM exam_subjects WHERE exam_id=? AND subject=?', (exam_id, subject)).fetchone():
                conn.execute('INSERT INTO exam_subjects(exam_id,subject,max_score,pass_score) VALUES(?,?,?,?)', (exam_id, subject, maximum, passing))
    return redirect(url_for('exam_config', exam_id=exam_id))


def build_summary(conn, exam):
    """Missing/absent/invalid scores do not participate in complete-score ranking."""
    from decimal import Decimal
    subjects = conn.execute('SELECT * FROM exam_subjects WHERE exam_id=? ORDER BY id', (exam['id'],)).fetchall()
    # Preserve legacy duplicates in storage, but count each subject only once.
    configs = {row['subject']: row for row in subjects}
    subjects = list(configs.values())
    students = conn.execute('SELECT * FROM students WHERE teacher_id=? ORDER BY id', (exam['teacher_id'],)).fetchall()
    rows = conn.execute('''SELECT s.*,es.max_score,es.pass_score FROM scores s JOIN students st ON st.id=s.student_id
        LEFT JOIN exam_subjects es ON es.id=(SELECT MAX(id) FROM exam_subjects WHERE exam_id=s.exam_id AND subject=s.subject)
        WHERE s.exam_id=? AND st.teacher_id=? AND s.id=(SELECT MAX(sc.id) FROM scores sc WHERE sc.student_id=s.student_id AND sc.exam_id=s.exam_id AND sc.subject=s.subject)''', (exam['id'], exam['teacher_id'])).fetchall()
    score_map = {(r['student_id'], r['subject']): normalise_score(r) for r in rows}
    ranking, incomplete, fails = [], [], []
    for student in students:
        total = Decimal('0')
        count = absent = invalid = fail_count = 0
        for subject in configs:
            row = score_map.get((student['id'], subject))
            if row is None:
                continue
            if not row['valid']:
                invalid += 1
            elif row['status'] == 'absent':
                absent += 1
            else:
                total += Decimal(str(row['score']))
                count += 1
                if row['fail']:
                    fail_count += 1
                    fails.append(dict(name=student['name'], subject=subject, score=row['score'], pass_score=row['pass_score']))
        result = dict(sid=student['id'], name=student['name'], total=float(total), avg=round(float(total) / count, 2) if count else None,
                      count=count, absent=absent, invalid=invalid, missing=len(configs)-count-absent-invalid, fail_count=fail_count,
                      complete=bool(configs) and count == len(configs), sort_total=total)
        (ranking if result['complete'] else incomplete).append(result)
    ranking.sort(key=lambda item: (-item['sort_total'], item['sid']))
    previous = None
    for pos, row in enumerate(ranking, start=1):
        row['rank'] = previous['rank'] if previous and row['sort_total'] == previous['sort_total'] else pos
        previous = row
    statistics = []
    for config in subjects:
        present = [score_map.get((s['id'], config['subject'])) for s in students]
        valid = [r for r in present if r and r['valid'] and r['status'] == 'recorded']
        values = [r['score'] for r in valid]
        passing = sum(not r['fail'] for r in valid)
        excellent = sum(r['score'] >= r['max_score'] * .85 for r in valid)
        statistics.append(dict(subject=config['subject'], count=len(values), expected=len(students), absent=sum(bool(r and r['valid'] and r['status'] == 'absent') for r in present),
                               average=round(sum(values)/len(values), 2) if values else None, maximum=max(values) if values else None, minimum=min(values) if values else None,
                               pass_rate=round(passing/len(values)*100, 1) if values else None, excellent_rate=round(excellent/len(values)*100, 1) if values else None))
    return dict(ranking=ranking, incomplete=incomplete, fails=fails, statistics=statistics, subjects=subjects,
                score_map=score_map, students=students, has_scores=any(r['count'] for r in ranking+incomplete),
                all_complete=bool(students) and not incomplete)


@app.route('/summary')
@login_required
def summary():
    with database() as conn:
        exams = own_exams(conn)
        eid = request.args.get('exam_id', type=int)
        if 'exam_id' in request.args and eid is None:
            abort(400, description='考试编号无效。')
        if eid is None and exams:
            eid = exams[0]['id']
        exam = require_exam(conn, eid) if eid is not None else None
        data = build_summary(conn, exam) if exam else dict(ranking=[], incomplete=[], fails=[], statistics=[], has_scores=False, all_complete=False)
    return render_template('summary.html', exam=exam, all_exams=exams, **data)


@app.route('/delete/<int:sid>', methods=['POST'])
@login_required
def delete_student(sid):
    with database(write=True) as conn:
        require_student(conn, sid)
        conn.execute('DELETE FROM scores WHERE student_id=?', (sid,))
        conn.execute('DELETE FROM students WHERE id=?', (sid,))
    return redirect(url_for('index'))


def require_score(conn, score_id):
    row = conn.execute('''SELECT s.*,st.name,st.teacher_id,e.name exam_name,es.max_score,es.pass_score FROM scores s
        JOIN students st ON st.id=s.student_id JOIN exams e ON e.id=s.exam_id
        LEFT JOIN exam_subjects es ON es.id=(SELECT MAX(id) FROM exam_subjects WHERE exam_id=s.exam_id AND subject=s.subject)
        WHERE s.id=?''', (score_id,)).fetchone()
    if not row:
        abort(404, description='没有这条成绩。')
    student = require_student(conn, row['student_id'])
    exam = require_exam(conn, row['exam_id'])
    require_pair(student, exam)
    return normalise_score(row)


@app.route('/edit_score/<int:score_id>', methods=['GET', 'POST'])
@login_required
def edit_score(score_id):
    error = None
    with database(write=request.method == 'POST') as conn:
        row = require_score(conn, score_id)
        if request.method == 'POST':
            try:
                if not row['config_valid']:
                    raise ValueError('此科目配置缺失或异常，请先到考试管理补齐配置，或删除此条异常成绩')
                value, status = parse_score(row)
            except ValueError as exc:
                error = str(exc)
            else:
                conn.execute('UPDATE scores SET score=?,status=? WHERE id=?', (value, status, score_id))
                return redirect(url_for('score_page', sid=row['student_id'], exam_id=row['exam_id']))
    return render_template('edit_score.html', rec=row, error=error)


@app.route('/delete_score/<int:score_id>', methods=['POST'])
@login_required
def delete_score(score_id):
    with database(write=True) as conn:
        row = require_score(conn, score_id)
        conn.execute('DELETE FROM scores WHERE student_id=? AND exam_id=? AND subject=?', (row['student_id'], row['exam_id'], row['subject']))
    flash('该科成绩已删除，现在为未录入状态。')
    return redirect(url_for('score_page', sid=row['student_id'], exam_id=row['exam_id']))


def sheet_title(name, used):
    name = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', name).strip().strip("'") or '考试'
    base = name[:31]
    title = base
    suffix = 1
    while title.casefold() in used:
        suffix += 1
        tag = f' ({suffix})'
        title = base[:31-len(tag)] + tag
    used.add(title.casefold())
    return title


def excel_value(cell, value):
    # All user strings must remain text even if they begin with '='.
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
        cell.value = value
        cell.data_type = 's'
    else:
        cell.value = value


@app.route('/export_excel')
@login_required
def export_excel():
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    head_fill = PatternFill('solid', fgColor='1E3A8A')
    red_fill = PatternFill('solid', fgColor='FECACA')
    with database() as conn:
        for exam in own_exams(conn):
            data = build_summary(conn, exam)
            ws = wb.create_sheet(sheet_title(exam['name'], used))
            headers = ['姓名'] + [s['subject'] for s in data['subjects']] + ['总分', '名次', '录入状态']
            for col, value in enumerate(headers, 1):
                cell = ws.cell(1, col)
                excel_value(cell, value)
                cell.fill = head_fill
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(horizontal='center')
            results = {r['sid']: r for r in data['ranking']+data['incomplete']}
            for idx, student in enumerate(data['students'], 2):
                excel_value(ws.cell(idx, 1), student['name'])
                for col, config in enumerate(data['subjects'], 2):
                    row = data['score_map'].get((student['id'], config['subject']))
                    value = '未录入' if row is None else row['score'] if row['valid'] and row['status']=='recorded' else row['display']
                    cell = ws.cell(idx, col)
                    excel_value(cell, value)
                    if row and row['fail']:
                        cell.fill = red_fill
                        cell.font = Font(color='DC2626', bold=True)
                result = results[student['id']]
                start = len(data['subjects']) + 2
                excel_value(ws.cell(idx, start), result['total'] if result['count'] else None)
                excel_value(ws.cell(idx, start+1), result.get('rank'))
                excel_value(ws.cell(idx, start+2), '已录全' if result['complete'] else f"未录入{result['missing']}科 / 缺考{result['absent']}科 / 异常{result['invalid']}科")
            ws.freeze_panes = 'B2'
            ws.auto_filter.ref = ws.dimensions
            for col in range(1, len(headers)+1):
                ws.column_dimensions[get_column_letter(col)].width = 16 if col==1 else 12
            ws.column_dimensions[get_column_letter(len(headers))].width = 40
    if not wb.worksheets:
        wb.create_sheet('全班成绩').append(['暂无考试数据'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='全班成绩.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=int(os.environ.get('PORT', 8080)), debug=False)
