"""Bounded Excel roster parsing; no database writes during preview."""
import io
import re
from zipfile import ZipFile

from openpyxl import load_workbook

MAX_STUDENTS = 500


def cell_text(cell, *, code=False):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type in ('f', 'e'):
        raise ValueError('不能使用公式或错误值，请粘贴为文本')
    if isinstance(value, str):
        result = value.strip()
    elif code and type(value) in (int, float) and value >= 0 and value == int(value):
        result = str(int(value))
        if len(result) > 15:
            raise ValueError('长学号请使用文本格式，避免 Excel 丢失精度')
        if re.fullmatch(r'0{1,50}', cell.number_format or ''):
            result = result.zfill(len(cell.number_format))
    else:
        raise ValueError('请使用文本格式' if not code else '学号请使用文本或非负整数')
    if re.search(r'[\x00-\x1f\x7f]', result):
        raise ValueError('不能包含换行或控制字符')
    return result


def parse_roster(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 100 or sum(e.file_size for e in entries) > 8 * 1024 * 1024:
                raise ValueError('Excel 内容过大，请拆分文件后重试')
            if any(e.flag_bits & 1 for e in entries):
                raise ValueError('请先取消 Excel 文件加密')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('无法读取文件，请上传未加密的 .xlsx 文件') from exc
    try:
        ws = wb.worksheets[0]
        # Do not trust worksheet dimensions supplied by the uploaded file.
        ws.reset_dimensions()
        iterator = ws.iter_rows(max_col=20)
        first = next(iterator, ())
        headers = [str(c.value).strip() if c.value is not None else '' for c in first]
        if headers.count('姓名') != 1 or headers.count('学号') > 1:
            raise ValueError('第一行需有且只有一列“姓名”，可选一列“学号”（位于前20列）')
        name_col = headers.index('姓名')
        code_col = headers.index('学号') if '学号' in headers else None
        rows = []
        for number, cells in enumerate(iterator, 2):
            if number > 2001:
                raise ValueError('表格行数过多，请去除多余空行或拆分文件')
            name_cell = cells[name_col]
            code_cell = cells[code_col] if code_col is not None else None
            if name_cell.value in (None, '') and (code_cell is None or code_cell.value in (None, '')):
                continue
            if len(rows) >= MAX_STUDENTS:
                raise ValueError('每次最多导入500名学生，请拆分文件')
            errors = []
            values = []
            for label, cell, is_code in [('姓名', name_cell, False), ('学号', code_cell, True)]:
                try:
                    values.append(cell_text(cell, code=is_code) if cell is not None else '')
                except (ValueError, OverflowError) as exc:
                    values.append('')
                    errors.append(f'{label}：{exc}')
            rows.append(dict(line=number, name=values[0], student_code=values[1], errors=errors))
        if not rows:
            raise ValueError('没有可导入的学生，请从第二行填写姓名和可选学号')
        return rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Excel 内容损坏，或单元格格式不受支持，请重新保存为 .xlsx') from exc
    finally:
        wb.close()


def validate_roster(rows, existing):
    names = {r['name'] for r in existing}
    codes = {r['student_code'] for r in existing if r['student_code']}
    seen_names, seen_codes = {}, {}
    result = []
    for raw in rows:
        row = {**raw, 'errors': list(raw.get('errors', []))}
        name, code = row['name'], row['student_code']
        if not 1 <= len(name) <= 100:
            row['errors'].append('姓名需为1至100位')
        if len(code) > 50:
            row['errors'].append('学号不能超过50位')
        if name in names:
            row['errors'].append('该班级已有同名学生')
        if code and code in codes:
            row['errors'].append('该班级已有相同学号')
        if name and name in seen_names:
            row['errors'].append(f'姓名与第{seen_names[name]}行重复')
        if code and code in seen_codes:
            row['errors'].append(f'学号与第{seen_codes[code]}行重复')
        if name:
            seen_names.setdefault(name, row['line'])
        if code:
            seen_codes.setdefault(code, row['line'])
        result.append(row)
    return result
